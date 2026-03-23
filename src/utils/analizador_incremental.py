# Analizador incremental para detectar cambios y sugerir mejoras
# Preserva trabajo manual de usuarios y solo actualiza lo necesario

import pandas as pd
import openpyxl
from pathlib import Path
from typing import Dict, List, Set, Optional
from datetime import datetime
import json

from .config import Config
from .text_processing import normalizar_texto
from .logger import ProjectLogger


class AnalizadorIncremental:
    """Analiza cambios incrementales en datos y sugiere mejoras de filtros"""
    
    def __init__(self, config: Optional[Config] = None):
        self.config = config or Config()
        self.logger = ProjectLogger('analizador_incremental', self.config.LOG_DIR)
        
    def analizar_cambios_taxonomia(self, datos_nuevos: pd.DataFrame) -> Dict:
        """Analiza cambios en taxonomía para sugerir filtros nuevos"""
        
        self.logger.info("🔍 Analizando cambios en taxonomía...")
        
        # Cargar taxonomía actual del PIVOT
        taxonomia_actual = self._cargar_taxonomia_pivot()
        
        # Extraeer taxonomía de datos nuevos
        taxonomia_nueva = self._extraer_taxonomia_datos(datos_nuevos)
        
        # Comparar y encontrar diferencias
        cambios = {
            'nuevos_nivel1': self._encontrar_nuevos_valores(taxonomia_actual['nivel1'], taxonomia_nueva['nivel1']),
            'nuevos_nivel2': self._encontrar_nuevos_valores(taxonomia_actual['nivel2'], taxonomia_nueva['nivel2']),
            'nuevos_nivel3': self._encontrar_nuevos_valores(taxonomia_actual['nivel3'], taxonomia_nueva['nivel3']),
            'nuevos_genericos': self._encontrar_nuevos_valores(taxonomia_actual['generico'], taxonomia_nueva['generico'])
        }
        
        # Generar sugerencias de filtros
        sugerencias = self._generar_sugerencias_filtros(cambios, datos_nuevos)
        
        self.logger.info(f"✅ Análisis completado: {sum(len(v) for v in cambios.values())} nuevos valores encontrados")
        
        return {
            'cambios_taxonomia': cambios,
            'sugerencias_filtros': sugerencias,
            'timestamp': datetime.now().isoformat()
        }
    
    def analizar_reporte_incremental(self, datos_nuevos: pd.DataFrame, archivo_reporte_anterior: Path = None) -> Dict:
        """Analiza qué licitaciones son realmente nuevas vs existentes"""
        
        self.logger.info("📊 Analizando reporte incremental...")
        
        # Buscar reporte anterior
        if not archivo_reporte_anterior:
            archivo_reporte_anterior = self._encontrar_reporte_mas_reciente()
        
        if not archivo_reporte_anterior or not archivo_reporte_anterior.exists():
            self.logger.info("📝 No hay reporte anterior - todos los datos son nuevos")
            return {
                'licitaciones_nuevas': datos_nuevos,
                'licitaciones_existentes': pd.DataFrame(),
                'licitaciones_vencidas': pd.DataFrame(),
                'estadisticas': {'nuevas': len(datos_nuevos), 'existentes': 0, 'vencidas': 0}
            }
        
        # Cargar reporte anterior
        reporte_anterior = self._cargar_reporte_anterior(archivo_reporte_anterior)
        
        # Comparar licitaciones
        resultado = self._comparar_licitaciones(datos_nuevos, reporte_anterior)
        
        self.logger.info(f"✅ Análisis incremental: {resultado['estadisticas']['nuevas']} nuevas, "
                        f"{resultado['estadisticas']['existentes']} existentes, "
                        f"{resultado['estadisticas']['vencidas']} vencidas")
        
        return resultado
    
    def generar_reporte_cambios(self, analisis_taxonomia: Dict, analisis_reporte: Dict) -> Dict:
        """Genera reporte completo de cambios y sugerencias"""
        
        reporte = {
            'timestamp': datetime.now().isoformat(),
            'resumen': {
                'taxonomia': {
                    'nuevos_nivel1': len(analisis_taxonomia['cambios_taxonomia']['nuevos_nivel1']),
                    'nuevos_nivel2': len(analisis_taxonomia['cambios_taxonomia']['nuevos_nivel2']),
                    'nuevos_nivel3': len(analisis_taxonomia['cambios_taxonomia']['nuevos_nivel3']),
                    'nuevos_genericos': len(analisis_taxonomia['cambios_taxonomia']['nuevos_genericos'])
                },
                'licitaciones': analisis_reporte['estadisticas'],
                'sugerencias_filtros': {
                    'inclusion': len(analisis_taxonomia['sugerencias_filtros']['inclusion']),
                    'exclusion': len(analisis_taxonomia['sugerencias_filtros']['exclusion'])
                }
            },
            'detalles': {
                'cambios_taxonomia': analisis_taxonomia['cambios_taxonomia'],
                'sugerencias_filtros': analisis_taxonomia['sugerencias_filtros']
            }
        }
        
        # Guardar reporte
        archivo_reporte = self.config.LOG_DIR / f"analisis_incremental_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
        with open(archivo_reporte, 'w', encoding='utf-8') as f:
            json.dump(reporte, f, indent=2, ensure_ascii=False)
        
        self.logger.info(f"📄 Reporte de cambios guardado: {archivo_reporte.name}")
        
        return reporte
    
    def _cargar_taxonomia_pivot(self) -> Dict[str, Set]:
        """Carga taxonomía actual del PIVOT_MAESTRO desde hoja 04-BASE (fuente canónica)."""
        taxonomia = {'nivel1': set(), 'nivel2': set(), 'nivel3': set(), 'generico': set()}

        try:
            wb = openpyxl.load_workbook(self.config.PIVOT_MAESTRO, data_only=True)

            # Determinar hoja a usar: 04-BASE es la fuente canónica; fallback a 01-AUDITORIA
            sheet_name = '04-BASE' if '04-BASE' in wb.sheetnames else '01-AUDITORIA'
            if sheet_name not in wb.sheetnames:
                self.logger.warning("⚠️ No se encontró hoja de taxonomía en PIVOT_MAESTRO (04-BASE / 01-AUDITORIA)")
                return taxonomia

            ws = wb[sheet_name]
            campo_idx = {}

            # Detectar fila de encabezado (máx 30 filas)
            for row in ws.iter_rows(max_row=30):
                vals = [normalizar_texto(str(c.value)) if c.value else '' for c in row[:10]]
                if any('NIVEL' in v and '1' in v for v in vals):
                    for j, cell in enumerate(row[:10]):
                        if not cell.value:
                            continue
                        norm = normalizar_texto(str(cell.value))
                        if 'NIVEL' in norm and '1' in norm:
                            campo_idx['nivel1'] = j
                        elif 'NIVEL' in norm and '2' in norm:
                            campo_idx['nivel2'] = j
                        elif 'NIVEL' in norm and '3' in norm:
                            campo_idx['nivel3'] = j
                        elif 'GENERICO' in norm:
                            campo_idx['generico'] = j
                    break

            if not campo_idx:
                self.logger.warning(f"⚠️ No se detectó encabezado de taxonomía en hoja {sheet_name}")
                return taxonomia

            header_row = next(
                (r for r in ws.iter_rows(max_row=30)
                 if any(normalizar_texto(str(c.value or '')) for c in r[:10])), None
            )
            start_row = (header_row[0].row + 1) if header_row else 2

            for row in ws.iter_rows(min_row=start_row):
                for campo, idx in campo_idx.items():
                    val = row[idx].value
                    if val and str(val).strip() not in ('', 'nan', 'None'):
                        taxonomia[campo].add(normalizar_texto(str(val)))

            wb.close()
            total = sum(len(v) for v in taxonomia.values())
            self.logger.info(f"[OK] Taxonomía cargada desde {sheet_name}: {total} valores")

        except Exception as e:
            self.logger.warning(f"⚠️ Error cargando taxonomía del PIVOT: {e}")

        return taxonomia
    
    def _extraer_taxonomia_datos(self, datos: pd.DataFrame) -> Dict[str, Set]:
        """Extrae taxonomía de datos nuevos"""
        
        taxonomia = {
            'nivel1': set(),
            'nivel2': set(),
            'nivel3': set(),
            'generico': set()
        }
        
        for col_name, tax_key in [('Nivel 1', 'nivel1'), ('Nivel 2', 'nivel2'), 
                                 ('Nivel 3', 'nivel3'), ('Genérico', 'generico')]:
            if col_name in datos.columns:
                valores = datos[col_name].dropna().astype(str)
                taxonomia[tax_key].update(normalizar_texto(v) for v in valores if v not in ['nan', 'None'])
        
        return taxonomia
    
    def _encontrar_nuevos_valores(self, valores_actuales: Set, valores_nuevos: Set) -> List[str]:
        """Encuentra valores que están en nuevos pero no en actuales"""
        return sorted(list(valores_nuevos - valores_actuales))
    
    def _generar_sugerencias_filtros(self, cambios: Dict, datos: pd.DataFrame) -> Dict:
        """Genera sugerencias de filtros basado en cambios detectados"""
        
        sugerencias = {
            'inclusion': [],
            'exclusion': []
        }
        
        # Analizar patrones en nuevos valores para sugerir filtros
        todos_nuevos = []
        for categoria, valores in cambios.items():
            todos_nuevos.extend(valores)
        
        # Palabras clave que sugieren inclusión (ingeniería/consultoría) — normalizadas sin tilde
        keywords_inclusion = ['CONSULTOR', 'INGENIR', 'DISENO', 'ARQUITECTUR', 'TECNIC', 'PROYECTO', 'DESARROLLO']

        # Palabras clave que sugieren exclusión — normalizadas sin tilde
        keywords_exclusion = ['SUMINISTRO', 'ARRIENDO', 'MANTENCI', 'LIMPIEZA', 'VIGILANC', 'ALIMENTA', 'TRANSPORT']

        for valor in todos_nuevos:
            valor_norm = normalizar_texto(valor)  # uppercase + sin tildes
            
            # Sugerir inclusión
            if any(kw in valor_norm for kw in keywords_inclusion):
                sugerencias['inclusion'].append({
                    'termino': valor,
                    'razon': 'Contiene palabras clave de ingeniería/consultoría',
                    'confianza': 'alta'
                })
            
            # Sugerir exclusión
            elif any(kw in valor_norm for kw in keywords_exclusion):
                sugerencias['exclusion'].append({
                    'termino': valor,
                    'razon': 'Contiene palabras clave de servicios operativos',
                    'confianza': 'alta'
                })
        
        return sugerencias
    
    def _encontrar_reporte_mas_reciente(self) -> Optional[Path]:
        """Encuentra el reporte más reciente en la carpeta de presentación"""
        
        archivos = list(self.config.PRESENTACION_DIR.glob("Reporte_Licitaciones_*.xlsx"))
        if not archivos:
            return None
        
        # Ordenar por fecha de modificación
        return max(archivos, key=lambda x: x.stat().st_mtime)
    
    def _cargar_reporte_anterior(self, archivo: Path) -> pd.DataFrame:
        """Carga reporte anterior preservando todas las columnas"""
        
        try:
            return pd.read_excel(archivo, sheet_name='Vigentes')
        except:
            try:
                return pd.read_excel(archivo)
            except Exception as e:
                self.logger.warning(f"⚠️ Error cargando reporte anterior: {e}")
                return pd.DataFrame()
    
    def _comparar_licitaciones(self, datos_nuevos: pd.DataFrame, reporte_anterior: pd.DataFrame) -> Dict:
        """Compara licitaciones nuevas vs existentes"""
        
        # Identificar columna de código único
        col_codigo = None
        for col in ['Numero Adquisición', 'Código Externo', 'CodigoExterno']:
            if col in datos_nuevos.columns:
                col_codigo = col
                break
        
        if not col_codigo:
            # Si no hay código, asumir que todas son nuevas
            return {
                'licitaciones_nuevas': datos_nuevos,
                'licitaciones_existentes': pd.DataFrame(),
                'licitaciones_vencidas': pd.DataFrame(),
                'estadisticas': {'nuevas': len(datos_nuevos), 'existentes': 0, 'vencidas': 0}
            }
        
        # Obtener códigos
        codigos_nuevos = set(datos_nuevos[col_codigo].astype(str))
        
        if col_codigo in reporte_anterior.columns:
            codigos_existentes = set(reporte_anterior[col_codigo].astype(str))
        else:
            codigos_existentes = set()
        
        # Clasificar licitaciones
        codigos_realmente_nuevos = codigos_nuevos - codigos_existentes
        codigos_ya_existentes = codigos_nuevos & codigos_existentes
        
        # Filtrar dataframes
        licitaciones_nuevas = datos_nuevos[datos_nuevos[col_codigo].astype(str).isin(codigos_realmente_nuevos)]
        licitaciones_existentes = reporte_anterior[reporte_anterior[col_codigo].astype(str).isin(codigos_ya_existentes)] if col_codigo in reporte_anterior.columns else pd.DataFrame()
        
        # Detectar licitaciones vencidas (en reporte anterior pero no en nuevos)
        codigos_vencidos = codigos_existentes - codigos_nuevos
        licitaciones_vencidas = reporte_anterior[reporte_anterior[col_codigo].astype(str).isin(codigos_vencidos)] if col_codigo in reporte_anterior.columns else pd.DataFrame()
        
        return {
            'licitaciones_nuevas': licitaciones_nuevas,
            'licitaciones_existentes': licitaciones_existentes,
            'licitaciones_vencidas': licitaciones_vencidas,
            'estadisticas': {
                'nuevas': len(licitaciones_nuevas),
                'existentes': len(licitaciones_existentes),
                'vencidas': len(licitaciones_vencidas)
            }
        }
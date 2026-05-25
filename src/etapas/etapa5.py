"""Etapa 5 — análisis incremental: combina licitaciones nuevas con reportes existentes preservando trabajo manual."""

import json
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook

from core.context import PipelineContext
from core.contracts import BaseStage, StageResult
from utils.analizador_incremental import AnalizadorIncremental
from utils.excel_formatter import guardar_formateado_reporte


class GeneradorReporteIncremental(BaseStage):
    """Combina el reporte de etapa 4 con reportes previos, añadiendo solo licitaciones nuevas y preservando el formateo manual."""
    
    # Nombres canónicos de columnas del reporte (generados por Etapa 4)
    NUMERO_ADQ_COL = 'Numero Adquisición'
    REGION_COL = 'Región'
    FECHA_CIERRE_COL = 'Fecha Cierre Licitación'

    def __init__(self) -> None:
        super().__init__()
        self.analizador = AnalizadorIncremental()

    def bind(self, context: PipelineContext) -> None:
        """Enlaza contexto/logger y redirige el logger del analizador al de este stage."""
        super().bind(context)
        self.analizador.logger = self.logger

    @property
    def name(self) -> str:
        return "incremental"
        
    def validate_inputs(self, context: PipelineContext) -> bool:
        permite_fallback = context.flags.get('allow_fallback', False)
        archivo_entrada = context.get_artifact('etapa4_output')
        
        if not archivo_entrada:
            if not permite_fallback:
                raise ValueError("Modo pipeline: Fallo al iniciar Etapa 5. Falta artefacto 'etapa4_output'.")
            
            # Buscar el archivo más reciente de etapa 4 en caso de fallback
            archivos = list(context.config.PRESENTACION_ORIGINAL_DIR.glob("Reporte_Licitaciones_*.xlsx"))
            if not archivos:
                raise FileNotFoundError(f"No hay archivo base de fallback en {context.config.PRESENTACION_ORIGINAL_DIR}")
            archivo_entrada = max(archivos, key=lambda f: f.stat().st_mtime)
            
        if not archivo_entrada.exists():
            raise FileNotFoundError(f"El archivo de entrada no existe físicamente: {archivo_entrada}")
            
        try:
            pd.read_excel(archivo_entrada, engine='openpyxl', nrows=1)
        except Exception as e:
            raise ValueError(f"El archivo base aportado a Etapa 5 es corrupto o no es Excel: {e!s}") from e
            
        return True

    def _execute(self, context: PipelineContext) -> StageResult:
        try:
            self.logger.section("ETAPA 5: REPORTE INCREMENTAL")
            inicio = datetime.now()
            
            self.dir_presentacion_incremental = self.config.PRESENTACION_INCREMENTAL_DIR
            self.dir_presentacion_incremental.mkdir(parents=True, exist_ok=True)
            
            archivo_reporte_nuevo = context.get_artifact('etapa4_output')
            if not archivo_reporte_nuevo:
                archivos = list(self.config.PRESENTACION_ORIGINAL_DIR.glob("Reporte_Licitaciones_*.xlsx"))
                archivo_reporte_nuevo = max(archivos, key=lambda f: f.stat().st_mtime)
                
            self.logger.info(f"📂 Cargando reporte nuevo: {archivo_reporte_nuevo.name}")
            datos_nuevos = pd.read_excel(archivo_reporte_nuevo)
            self.logger.info(f"📊 Datos cargados: {len(datos_nuevos)} licitaciones")
            
            # 2. Buscar reporte incremental anterior
            reporte_anterior = self._encontrar_reporte_incremental_anterior()
            
            # Pre-computar análisis incremental UNA VEZ — evita doble lectura del archivo anterior
            analisis_reporte = (
                self.analizador.analizar_reporte_incremental(datos_nuevos, reporte_anterior)
                if reporte_anterior else None
            )

            if reporte_anterior:
                self.logger.info(f"📋 Reporte anterior encontrado: {reporte_anterior.name}")
                resultado = self._actualizar_reporte_existente(datos_nuevos, reporte_anterior, analisis_reporte)
            else:
                self.logger.info("📄 No hay reporte anterior, creando reporte inicial")
                resultado = self._crear_reporte_inicial(datos_nuevos)
            
            # 3. Generar análisis de cambios (análisis ya computado, sin releer disco)
            analisis = self._generar_analisis_cambios(datos_nuevos, reporte_anterior, analisis_reporte)
            
            # 4. Guardar sugerencias para PIVOT
            self._guardar_sugerencias_pivot(analisis)
            
            
            rutas_generadas = []
            if resultado.get('archivo_generado'):
                rutas_generadas.append(resultado['archivo_generado'])
                context.add_artifact('etapa5_output', resultado['archivo_generado'])
            
            tiempo_total = datetime.now() - inicio
            resultado['tiempo_ejecucion'] = str(tiempo_total)
            
            # Serializar Path → str para compatibilidad con json.dumps en observabilidad
            resultado_serial = {
                k: str(v) if isinstance(v, Path) else v
                for k, v in resultado.items()
            }
            
            self.logger.info(f"✅ Etapa 5 completada en {tiempo_total}")
            return StageResult(
                success=True,
                stage_name=self.name,
                files_produced=rutas_generadas,
                metrics_produced=resultado_serial,
                custom_data={'detalles': resultado_serial}
            )
            
        except Exception as e:
            self.logger.error(f"❌ Error crítico en Etapa 5: {e}", exc_info=True)
            return StageResult(success=False, stage_name=self.name, error_message=str(e))
    
    def _encontrar_reporte_incremental_anterior(self) -> Path | None:
        """Encuentra el reporte incremental más reciente"""
        
        archivos = list(self.dir_presentacion_incremental.glob("Reporte_Incremental_*.xlsx"))
        
        if archivos:
            # Ordenar por fecha de modificación
            archivo_mas_reciente = max(archivos, key=lambda x: x.stat().st_mtime)
            return archivo_mas_reciente
        
        return None
    
    def _actualizar_reporte_existente(self, datos_nuevos: pd.DataFrame, reporte_anterior: Path, analisis_reporte: dict | None = None) -> dict:
        """Actualiza reporte existente preservando trabajo manual a nivel de celda.

        Copia el archivo anterior (conservando colores y formatos manuales), luego:
        - Actualiza 'Días para cierre' de filas existentes
        - Mueve vencidas a hoja 'Vencidas'
        - Agrega filas nuevas al final
        """
        self.logger.info("🔄 Actualizando reporte existente (preservando formato manual)...")

        # Usar análisis pre-computado si se provee (evita doble lectura de disco desde _execute)
        if analisis_reporte is None:
            analisis_reporte = self.analizador.analizar_reporte_incremental(datos_nuevos, reporte_anterior)
        licitaciones_nuevas = analisis_reporte['licitaciones_nuevas']
        licitaciones_existentes = analisis_reporte['licitaciones_existentes']
        licitaciones_vencidas = analisis_reporte['licitaciones_vencidas']

        self.logger.info(f"📊 Análisis: {len(licitaciones_nuevas)} nuevas, "
                         f"{len(licitaciones_existentes)} existentes, "
                         f"{len(licitaciones_vencidas)} vencidas")

        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        archivo_incremental = self.dir_presentacion_incremental / f"Reporte_Incremental_{timestamp}.xlsx"

        # Copiar archivo anterior para preservar TODOS los formatos manuales (colores, filtros, etc.)
        shutil.copy2(reporte_anterior, archivo_incremental)

        wb = load_workbook(archivo_incremental)
        ws = wb['Vigentes'] if 'Vigentes' in wb.sheetnames else wb.active

        # Construir mapa: nombre_columna -> índice 1-based
        header_map = {
            str(ws.cell(1, c).value): c
            for c in range(1, ws.max_column + 1)
            if ws.cell(1, c).value is not None
        }
        col_codigo = header_map.get(self.NUMERO_ADQ_COL)
        col_dias = header_map.get('Días para cierre')

        # 1. Actualizar 'Días para cierre' de licitaciones existentes sin tocar su formato
        if col_codigo and col_dias and not datos_nuevos.empty:
            col_codigo_df = next(
                (c for c in [self.NUMERO_ADQ_COL, 'Código Externo', 'CodigoExterno']
                 if c in datos_nuevos.columns), None
            )
            if col_codigo_df and 'Días para cierre' in datos_nuevos.columns:
                dias_map = dict(zip(
                    datos_nuevos[col_codigo_df].astype(str),
                    datos_nuevos['Días para cierre'],
                    strict=False,
                ))
                for row in ws.iter_rows(min_row=2):
                    codigo_val = row[col_codigo - 1].value
                    if codigo_val and str(codigo_val) in dias_map:
                        row[col_dias - 1].value = dias_map[str(codigo_val)]

        # 2. Mover vencidas a hoja separada
        if not licitaciones_vencidas.empty and col_codigo:
            self._mover_licitaciones_vencidas(wb, ws, licitaciones_vencidas, col_codigo)

        # 3. Agregar licitaciones nuevas al final preservando estructura de columnas
        if not licitaciones_nuevas.empty:
            cols_ws = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            for _, row_data in licitaciones_nuevas.iterrows():
                last_row = ws.max_row + 1
                for col_idx, col_name in enumerate(cols_ws, 1):
                    if col_name and col_name in row_data.index:
                        ws.cell(last_row, col_idx).value = row_data.get(col_name, '')
            self.logger.info(f"➕ {len(licitaciones_nuevas)} licitaciones nuevas agregadas")

        wb.save(archivo_incremental)
        self.logger.info(f"💾 Reporte incremental guardado: {archivo_incremental.name}")

        return {
            'archivo_generado': archivo_incremental,
            'licitaciones_nuevas': len(licitaciones_nuevas),
            'licitaciones_existentes': len(licitaciones_existentes),
            'licitaciones_vencidas': len(licitaciones_vencidas),
            'tipo': 'incremental'
        }
    
    def _crear_reporte_inicial(self, datos_nuevos: pd.DataFrame) -> dict:
        """Crea el primer reporte incremental"""
        
        self.logger.info("📄 Creando reporte incremental inicial...")
        
        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        archivo_incremental = self.dir_presentacion_incremental / f"Reporte_Incremental_{timestamp}.xlsx"
        
        # Crear Excel con formato mejorado
        guardar_formateado_reporte(datos_nuevos, archivo_incremental, "Vigentes")
        self.logger.info(f"💾 Reporte inicial guardado: {archivo_incremental.name}")
        
        return {
            'archivo_generado': archivo_incremental,
            'licitaciones_nuevas': len(datos_nuevos),
            'licitaciones_existentes': 0,
            'licitaciones_vencidas': 0,
            'tipo': 'inicial'
        }

    def _mover_licitaciones_vencidas(
        self, wb: Workbook, ws_vigentes: Any, vencidas: pd.DataFrame, col_codigo_idx: int
    ) -> None:
        """Copia licitaciones vencidas a hoja 'Vencidas' y las elimina de 'Vigentes'."""
        if 'Vencidas' not in wb.sheetnames:
            wb.create_sheet('Vencidas')
        ws_vencidas = wb['Vencidas']

        # Copiar encabezado si la hoja está vacía (max_row puede ser None o 1 en openpyxl)
        if (ws_vencidas.max_row or 0) <= 1 and ws_vencidas.cell(1, 1).value is None:
            for c in range(1, ws_vigentes.max_column + 1):
                ws_vencidas.cell(1, c).value = ws_vigentes.cell(1, c).value

        col_codigo_name = next(
            (c for c in [self.NUMERO_ADQ_COL, 'Código Externo', 'CodigoExterno']
             if c in vencidas.columns), None
        )
        if not col_codigo_name:
            self.logger.warning("⚠️ No se identificó columna de código para mover vencidas")
            return

        codigos_vencidos = set(vencidas[col_codigo_name].astype(str))
        rows_to_delete = []

        for row in ws_vigentes.iter_rows(min_row=2):
            val = row[col_codigo_idx - 1].value
            if val and str(val) in codigos_vencidos:
                new_row = ws_vencidas.max_row + 1
                for col_idx, cell in enumerate(row, 1):
                    ws_vencidas.cell(new_row, col_idx).value = cell.value
                rows_to_delete.append(row[0].row)

        for row_idx in reversed(rows_to_delete):
            ws_vigentes.delete_rows(row_idx)

        self.logger.info(f"📦 {len(rows_to_delete)} licitaciones movidas a hoja 'Vencidas'")
    
    def _generar_analisis_cambios(self, datos_nuevos: pd.DataFrame, reporte_anterior: Path | None, analisis_reporte: dict | None = None) -> dict:
        """Genera análisis completo de cambios."""
        analisis_taxonomia = self.analizador.analizar_cambios_taxonomia(datos_nuevos)
        
        if reporte_anterior and analisis_reporte is None:
            # Computar solo si no se proveyó un análisis pre-calculado
            analisis_reporte = self.analizador.analizar_reporte_incremental(datos_nuevos, reporte_anterior)
        elif not reporte_anterior:
            analisis_reporte = {
                'licitaciones_nuevas': datos_nuevos,
                'licitaciones_existentes': pd.DataFrame(),
                'licitaciones_vencidas': pd.DataFrame(),
                'estadisticas': {
                    'nuevas': len(datos_nuevos),
                    'existentes': 0,
                    'vencidas': 0
                }
            }
        
        return self.analizador.generar_reporte_cambios(analisis_taxonomia, analisis_reporte)
    
    def _guardar_sugerencias_pivot(self, analisis: dict) -> None:
        """Guarda sugerencias para actualizar PIVOT_MAESTRO"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        archivo_sugerencias = self.config.LOG_DIR / f"sugerencias_pivot_etapa5_{timestamp}.json"

        def _serializar(obj: Any) -> Any:
            if isinstance(obj, pd.DataFrame):
                return obj.to_dict(orient='records')
            if isinstance(obj, set):
                return sorted(obj)
            return str(obj)

        with archivo_sugerencias.open('w', encoding='utf-8') as f:
            json.dump(analisis, f, indent=2, ensure_ascii=False, default=_serializar)
        
        self.logger.info(f"💡 Sugerencias PIVOT guardadas: {archivo_sugerencias.name}")


def main() -> int:
    """Función principal para ejecutar Etapa 5 de forma independiente."""
    from utils.config import Config
    from utils.logger import configurar_consola_utf8
    configurar_consola_utf8()

    context = PipelineContext(config=Config())
    context.flags['allow_fallback'] = True
    gen = GeneradorReporteIncremental()
    resultado = gen.run(context)

    if resultado.success:
        detalles = resultado.custom_data.get('detalles', {})
        print("\n✅ ETAPA 5 COMPLETADA")
        print(f"   * Tipo:                  {detalles.get('tipo', 'N/A')}")
        print(f"   * Licitaciones nuevas:   {detalles.get('licitaciones_nuevas', 0)}")
        print(f"   * Licitaciones existentes: {detalles.get('licitaciones_existentes', 0)}")
        print(f"   * Licitaciones vencidas: {detalles.get('licitaciones_vencidas', 0)}")
        return 0

    print(f"\n❌ ERROR: {resultado.error_message}")
    return 1


if __name__ == "__main__":
    exit(main())


"""ProfileRegistry y ProfileLoader: cargan equipos y sus reglas desde PIVOT_MAESTRO.xlsx.

Responsabilidades separadas:
    - ProfileRegistry  : lee la hoja "00-Equipos" → lista de EquipoInfo.
    - ProfileLoader    : lee la hoja específica de un equipo → FilterProfile.

Ambos son inyectables (acepta ruta de PIVOT) → testeables sin tocar producción.
"""
from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING, Any

import openpyxl
import pandas as pd

from core.filter_profile import (
    EquipoInfo,
    FilterProfile,
    IntencionGlobal,
)
from utils.text_processing import normalizar_texto

if TYPE_CHECKING:
    from collections.abc import Mapping, Sequence

# ---------------------------------------------------------------------------
# Constantes de layout del PIVOT (única fuente de verdad)
# ---------------------------------------------------------------------------

#: Hoja con el catálogo de equipos.
HOJA_EQUIPOS: str = "00-Equipos"

#: Hoja con la invariante de negocio MP (intent gate). OPCIONAL: si no existe,
#: el filtro de intención queda deshabilitado y el sistema se comporta como antes.
HOJA_INTENCION_GLOBAL: str = "01-Intencion_Global"

#: Nombres esperados de las columnas en HOJA_INTENCION_GLOBAL (case-insensitive).
COL_INTENCION_REQUERIDA: str = "intencion_requerida"
COL_INTENCION_VETADA: str = "intencion_vetada"

#: Fila (1-indexed openpyxl) donde empiezan los datos en cualquier hoja de filtros.
FILTROS_DATA_START_ROW: int = 6

#: Fila (0-indexed pandas) del header de las hojas de filtros.
FILTROS_HEADER_ROW: int = 4

#: Mapeo columna (0-indexed) → clave del FilterProfile.incluir / .excluir.
COLUMNAS_INCLUIR: Mapping[int, str] = {0: "nombre", 1: "nivel1", 2: "nivel2", 3: "nivel3"}
COLUMNAS_EXCLUIR: Mapping[int, str] = {
    4: "nombre", 5: "nivel1", 6: "nivel2", 7: "nivel3",
    8: "generico", 9: "componente", 10: "organismo", 11: "valor",
}
COLUMNA_BYPASS: int = 12
COLUMNA_EXCLUSION_DURA: int = 13


# ---------------------------------------------------------------------------
# Excepciones de dominio
# ---------------------------------------------------------------------------

class PivotConfigError(Exception):
    """Error de configuración del PIVOT_MAESTRO. Mensaje legible para el usuario."""


# ---------------------------------------------------------------------------
# Registro de equipos
# ---------------------------------------------------------------------------

class ProfileRegistry:
    """Catálogo de equipos disponibles. Lee la hoja 00-Equipos del PIVOT.

    Layout esperado de la hoja 00-Equipos (header en fila 1):
        | codigo | nombre | hoja_filtros | descripcion | activo |
        | TELECOM| Telecomunicaciones | 06-Telecom | ... | TRUE |
        | ARQ    | Arquitectura       | 06-Arquitectura | ... | TRUE |
    """

    def __init__(self, pivot_path: Path) -> None:
        self._pivot_path = Path(pivot_path)

    def listar_equipos(self) -> list[EquipoInfo]:
        """Devuelve todos los equipos definidos en el catálogo (activos e inactivos)."""
        if not self._pivot_path.exists():
            raise PivotConfigError(f"PIVOT_MAESTRO no encontrado: {self._pivot_path}")

        try:
            df = pd.read_excel(self._pivot_path, sheet_name=HOJA_EQUIPOS, dtype=str)
        except ValueError as e:
            # openpyxl/pandas error cuando la hoja no existe
            raise PivotConfigError(
                f"Falta la hoja '{HOJA_EQUIPOS}' en {self._pivot_path.name}. "
                f"Esta hoja define qué equipos existen."
            ) from e

        cols_requeridas = {"codigo", "nombre", "hoja_filtros"}
        cols_presentes = {str(c).strip().lower() for c in df.columns}
        faltantes = cols_requeridas - cols_presentes
        if faltantes:
            raise PivotConfigError(
                f"Hoja '{HOJA_EQUIPOS}' carece de columnas: {sorted(faltantes)}. "
                f"Encontradas: {sorted(cols_presentes)}"
            )

        # Normalizar nombres de columnas (case-insensitive)
        df.columns = [str(c).strip().lower() for c in df.columns]

        equipos: list[EquipoInfo] = []
        for _, row in df.iterrows():
            codigo = str(row.get("codigo") or "").strip()
            if not codigo or codigo.lower() in ("nan", "none"):
                continue
            activo_raw = str(row.get("activo") or "true").strip().lower()
            activo = activo_raw not in ("false", "0", "no", "")
            equipos.append(EquipoInfo(
                codigo=codigo.upper(),
                nombre=str(row.get("nombre") or codigo).strip(),
                hoja_filtros=str(row.get("hoja_filtros") or "").strip(),
                descripcion=str(row.get("descripcion") or "").strip() if "descripcion" in df.columns else "",
                activo=activo,
            ))
        return equipos

    def equipos_activos(self) -> list[EquipoInfo]:
        """Solo los equipos marcados como activos. Usado por el selector del GUI."""
        return [e for e in self.listar_equipos() if e.activo]

    def buscar(self, codigo: str) -> EquipoInfo:
        """Devuelve el EquipoInfo del código dado. Lanza PivotConfigError si no existe."""
        codigo_norm = (codigo or "").strip().upper()
        for e in self.listar_equipos():
            if e.codigo == codigo_norm:
                return e
        disponibles = [e.codigo for e in self.listar_equipos()]
        raise PivotConfigError(
            f"Equipo '{codigo}' no existe en el catálogo. Disponibles: {disponibles}"
        )


# ---------------------------------------------------------------------------
# Cargador de perfiles
# ---------------------------------------------------------------------------

class ProfileLoader:
    """Carga el FilterProfile de un equipo desde su hoja específica del PIVOT."""

    def __init__(self, pivot_path: Path) -> None:
        self._pivot_path = Path(pivot_path)
        # Cache de IntencionGlobal: se lee una sola vez por instancia (es invariante).
        self._intencion_cache: IntencionGlobal | None = None

    def cargar_intencion_global(self) -> IntencionGlobal:
        """Carga la hoja `01-Intencion_Global` con las palabras de intent gate.

        Si la hoja no existe o está vacía, devuelve un IntencionGlobal vacío
        (feature deshabilitado, comportamiento legacy). Cached por instancia.
        """
        if self._intencion_cache is not None:
            return self._intencion_cache

        if not self._pivot_path.exists():
            self._intencion_cache = IntencionGlobal()
            return self._intencion_cache

        try:
            df = pd.read_excel(self._pivot_path, sheet_name=HOJA_INTENCION_GLOBAL, dtype=str)
        except ValueError:
            # Hoja no existe → feature deshabilitado (backward compat)
            self._intencion_cache = IntencionGlobal()
            return self._intencion_cache

        # Normalizar nombres de columnas (case-insensitive, espacios)
        df.columns = [str(c).strip().lower().replace(" ", "_") for c in df.columns]

        def leer_col(nombre: str) -> tuple[str, ...]:
            if nombre not in df.columns:
                return ()
            seen: set[str] = set()
            out: list[str] = []
            for v in df[nombre].dropna().astype(str):
                if v.lower() in ("nan", "none", ""):
                    continue
                norm = normalizar_texto(v)
                if not norm or norm in seen:
                    continue
                seen.add(norm)
                out.append(norm)
            return tuple(out)

        self._intencion_cache = IntencionGlobal(
            requerida=leer_col(COL_INTENCION_REQUERIDA),
            vetada=leer_col(COL_INTENCION_VETADA),
        )
        return self._intencion_cache

    def cargar(self, equipo: EquipoInfo) -> FilterProfile:
        """Lee la hoja del equipo y construye su FilterProfile.

        Raises:
            PivotConfigError: Si la hoja no existe o tiene estructura inválida.
        """
        if not self._pivot_path.exists():
            raise PivotConfigError(f"PIVOT_MAESTRO no encontrado: {self._pivot_path}")

        try:
            df = pd.read_excel(
                self._pivot_path,
                sheet_name=equipo.hoja_filtros,
                header=FILTROS_HEADER_ROW,
            )
        except ValueError as e:
            raise PivotConfigError(
                f"Equipo '{equipo.codigo}' apunta a hoja '{equipo.hoja_filtros}' "
                f"pero esa hoja no existe en {self._pivot_path.name}"
            ) from e

        def leer_columna(col_idx: int) -> tuple[str, ...]:
            """Lee una columna por índice, normaliza, descarta vacíos y duplicados."""
            if col_idx >= len(df.columns):
                return ()
            seen: set[str] = set()
            resultado: list[str] = []
            for v in df.iloc[:, col_idx].dropna().astype(str):
                if v.lower() in ("nan", "none", ""):
                    continue
                norm = normalizar_texto(v)
                if not norm or norm in seen:
                    continue
                seen.add(norm)
                resultado.append(norm)
            return tuple(resultado)

        incluir: dict[str, Sequence[str]] = {
            key: leer_columna(idx) for idx, key in COLUMNAS_INCLUIR.items()
        }
        excluir: dict[str, Sequence[str]] = {
            key: leer_columna(idx) for idx, key in COLUMNAS_EXCLUIR.items()
        }
        bypass = leer_columna(COLUMNA_BYPASS)
        exclusion_dura = leer_columna(COLUMNA_EXCLUSION_DURA)

        return FilterProfile(
            equipo=equipo,
            incluir=incluir,
            excluir=excluir,
            bypass=bypass,
            exclusion_dura=exclusion_dura,
            intencion_global=self.cargar_intencion_global(),
        )


# ---------------------------------------------------------------------------
# Validador (chequeos de calidad sobre un perfil cargado)
# ---------------------------------------------------------------------------

class ProfileValidator:
    """Valida un FilterProfile y devuelve lista de problemas (vacía si todo OK).

    NO levanta excepciones — devuelve diagnósticos para que el GUI los muestre.
    """

    @staticmethod
    def validar(profile: FilterProfile) -> list[str]:
        problemas: list[str] = []

        if profile.esta_vacio():
            problemas.append(
                f"El equipo '{profile.equipo.codigo}' no tiene reglas de inclusión. "
                f"El filtrado no producirá resultados."
            )

        # Cada keyword de inclusión debe tener al menos 2 caracteres
        for key, palabras in profile.incluir.items():
            cortas = [p for p in palabras if len(p) < 2]
            if cortas:
                problemas.append(
                    f"Inclusión[{key}] contiene palabras de 1 carácter: {cortas}. "
                    f"Eliminarlas — generan ruido masivo."
                )

        # Bypass también debe respetar longitud mínima
        bypass_cortas = [p for p in profile.bypass if len(p) < 2]
        if bypass_cortas:
            problemas.append(
                f"Bypass contiene palabras muy cortas: {bypass_cortas}. "
                f"Riesgo de matches espurios."
            )

        return problemas


# ---------------------------------------------------------------------------
# Onboarding: copiar reglas de un equipo a otro
# ---------------------------------------------------------------------------

#: Nº de columnas de datos en cada hoja de filtros (A..N → 14 columnas).
_TOTAL_COLUMNAS_FILTROS: int = max(
    max(COLUMNAS_INCLUIR.keys()),
    max(COLUMNAS_EXCLUIR.keys()),
    COLUMNA_BYPASS,
    COLUMNA_EXCLUSION_DURA,
) + 1


class ProfileCopier:
    """Copia las reglas de filtrado de una hoja-equipo a otra dentro del mismo PIVOT.

    Útil para onboarding: un equipo nuevo (p. ej. ARQ) puede partir de la
    configuración de TELECOM y luego ajustar. Solo se tocan las celdas de
    datos (filas >= FILTROS_DATA_START_ROW, columnas 0.._TOTAL_COLUMNAS_FILTROS-1);
    headers y formato existentes en destino se preservan.
    """

    def __init__(self, pivot_path: Path) -> None:
        self._pivot_path = Path(pivot_path)

    def copiar(self, hoja_origen: str, hoja_destino: str, *, sobrescribir: bool = False) -> int:
        """Copia las reglas de `hoja_origen` a `hoja_destino`.

        Args:
            sobrescribir: si False y la hoja destino ya tiene datos, levanta
                PivotConfigError (protección contra borrados accidentales).

        Returns:
            Número de filas de datos copiadas (excluyendo filas totalmente vacías).
        """
        if hoja_origen == hoja_destino:
            raise PivotConfigError("La hoja origen y destino no pueden ser la misma.")
        if not self._pivot_path.exists():
            raise PivotConfigError(f"PIVOT_MAESTRO no encontrado: {self._pivot_path}")

        wb = openpyxl.load_workbook(self._pivot_path)
        try:
            if hoja_origen not in wb.sheetnames:
                raise PivotConfigError(f"Hoja origen no existe: '{hoja_origen}'")
            if hoja_destino not in wb.sheetnames:
                raise PivotConfigError(f"Hoja destino no existe: '{hoja_destino}'")

            ws_dst = wb[hoja_destino]
            if not sobrescribir and self._tiene_datos(ws_dst):
                raise PivotConfigError(
                    f"La hoja destino '{hoja_destino}' ya contiene reglas. "
                    f"Usa sobrescribir=True para reemplazarlas."
                )

            ws_src = wb[hoja_origen]
            self._limpiar_datos(ws_dst)
            filas_copiadas = self._copiar_datos(ws_src, ws_dst)
            wb.save(self._pivot_path)
            return filas_copiadas
        finally:
            wb.close()

    def _tiene_datos(self, ws: Any) -> bool:
        for row in ws.iter_rows(
            min_row=FILTROS_DATA_START_ROW,
            max_col=_TOTAL_COLUMNAS_FILTROS,
            values_only=True,
        ):
            if any(v is not None and str(v).strip() != "" for v in row):
                return True
        return False

    def _limpiar_datos(self, ws: Any) -> None:
        if ws.max_row < FILTROS_DATA_START_ROW:
            return
        for row in ws.iter_rows(
            min_row=FILTROS_DATA_START_ROW,
            max_row=ws.max_row,
            max_col=_TOTAL_COLUMNAS_FILTROS,
        ):
            for cell in row:
                cell.value = None

    def _copiar_datos(self, ws_src: Any, ws_dst: Any) -> int:
        filas = 0
        for offset, row in enumerate(
            ws_src.iter_rows(
                min_row=FILTROS_DATA_START_ROW,
                max_col=_TOTAL_COLUMNAS_FILTROS,
                values_only=True,
            ),
            start=0,
        ):
            if not any(v is not None and str(v).strip() != "" for v in row):
                continue
            for col_idx, value in enumerate(row, start=1):
                ws_dst.cell(row=FILTROS_DATA_START_ROW + offset, column=col_idx, value=value)
            filas += 1
        return filas

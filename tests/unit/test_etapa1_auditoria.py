"""Tests unitarios para AuditorTaxonomia (Etapa 1).

Cubre:
- validate_inputs: con y sin artefacto, con archivo en disco
- _cargar_valores_pivot: happy path, sin header, celdas vacías
- _procesar_campos: nuevos, conocidos, umbral, columna faltante, similares
- _generar_reporte: creación de Excel con sheets correctas, ordenamiento
- _cargar_parametros_pivot: carga y fallo gracioso
- run(): flujo completo y fallo por PIVOT ausente
"""
import pytest
import pandas as pd
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_pivot_xlsx(path: Path) -> None:
    """Crea PIVOT_MAESTRO.xlsx mínimo con hoja 04-BASE."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "04-BASE"
    ws.append(["Nivel 1", "Nivel 2", "Nivel 3", "Genérico"])
    ws.append(["Servicios", "Tecnología", "Consultoría", "Software"])
    ws.append(["Obras", "Construcción", "Edificación", "Civil"])
    wb.save(path)


def _make_licitaciones_xlsx(path: Path) -> None:
    """Crea Excel de licitaciones con valores conocidos en pivot."""
    pd.DataFrame({
        "Nivel 1": ["Servicios", "Obras"],
        "Nivel 2": ["Tecnología", "Construcción"],
        "Nivel 3": ["Consultoría", "Edificación"],
        "Genérico": ["Software", "Civil"],
    }).to_excel(path, index=False, engine="openpyxl")


def _stage_con_config(tmp_path, logger_mock, *, pivot=True, hallazgos=True,
                      input_dir=None):
    """Crea AuditorTaxonomia con Config apuntando a tmp_path."""
    from utils.config import Config
    from core.context import PipelineContext
    from etapas.etapa1 import AuditorTaxonomia

    kwargs = {"env": "testing", "PIVOT_DIR": tmp_path}
    if hallazgos:
        hdir = tmp_path / "hallazgos"
        hdir.mkdir(exist_ok=True)
        kwargs["HALLAZGOS_DIR"] = hdir
    if input_dir:
        kwargs["INPUT_DIR"] = input_dir

    if pivot:
        _make_pivot_xlsx(tmp_path / "PIVOT_MAESTRO.xlsx")

    config = Config(**kwargs)
    stage = AuditorTaxonomia()
    stage._context = PipelineContext(config=config)
    stage._logger = logger_mock
    return stage


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def etapa1(config, logger_mock):
    """AuditorTaxonomia básico con config de test global."""
    from etapas.etapa1 import AuditorTaxonomia
    from core.context import PipelineContext
    stage = AuditorTaxonomia()
    stage._context = PipelineContext(config=config)
    stage._logger = logger_mock
    return stage


# ---------------------------------------------------------------------------
# Tests: validate_inputs
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestValidateInputs:

    def test_retorna_false_sin_artefacto_ni_archivo(self, tmp_path, logger_mock):
        from utils.config import Config
        from core.context import PipelineContext
        from etapas.etapa1 import AuditorTaxonomia

        config = Config(env="testing",
                        INPUT_DIR=tmp_path / "nonexistent_input")
        ctx = PipelineContext(config=config)
        stage = AuditorTaxonomia()
        stage._context = ctx
        stage._logger = logger_mock

        assert stage.validate_inputs(ctx) is False

    def test_retorna_true_con_artefacto(self, etapa1, context, tmp_path):
        archivo = tmp_path / "licitaciones.xlsx"
        archivo.touch()
        context.add_artifact("etapa0_output", archivo)
        assert etapa1.validate_inputs(context) is True

    def test_retorna_true_con_archivo_en_disco(self, tmp_path, logger_mock):
        from utils.config import Config
        from core.context import PipelineContext
        from etapas.etapa1 import AuditorTaxonomia

        input_dir = tmp_path / "INPUT"
        input_dir.mkdir()
        licitaciones = input_dir / "Licitacion_Publicada.xlsx"
        licitaciones.touch()

        config = Config(env="testing", INPUT_DIR=input_dir)
        ctx = PipelineContext(config=config)
        stage = AuditorTaxonomia()
        stage._context = ctx
        stage._logger = logger_mock

        assert stage.validate_inputs(ctx) is True


# ---------------------------------------------------------------------------
# Tests: _cargar_valores_pivot
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestCargarValoresPivot:

    def test_happy_path_retorna_cuatro_campos(self, tmp_path, logger_mock):
        stage = _stage_con_config(tmp_path, logger_mock)
        valores = stage._cargar_valores_pivot()

        assert isinstance(valores, dict)
        assert set(valores.keys()) == {"Nivel 1", "Nivel 2", "Nivel 3", "Generico"}

    def test_valores_estan_normalizados(self, tmp_path, logger_mock):
        from utils import normalizar_texto
        stage = _stage_con_config(tmp_path, logger_mock)
        valores = stage._cargar_valores_pivot()

        assert normalizar_texto("Servicios") in valores["Nivel 1"]
        assert normalizar_texto("Obras") in valores["Nivel 1"]
        assert normalizar_texto("Tecnología") in valores["Nivel 2"]
        assert normalizar_texto("Software") in valores["Generico"]

    def test_sin_header_lanza_value_error(self, tmp_path, logger_mock):
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "04-BASE"
        ws.append(["Col A", "Col B"])  # Sin encabezado válido NIVEL
        wb.save(tmp_path / "PIVOT_MAESTRO.xlsx")

        from utils.config import Config
        from core.context import PipelineContext
        from etapas.etapa1 import AuditorTaxonomia

        config = Config(env="testing", PIVOT_DIR=tmp_path)
        stage = AuditorTaxonomia()
        stage._context = PipelineContext(config=config)
        stage._logger = logger_mock

        with pytest.raises(ValueError, match="No se encontró header"):
            stage._cargar_valores_pivot()

    def test_ignora_celdas_vacias_y_none(self, tmp_path, logger_mock):
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "04-BASE"
        ws.append(["Nivel 1", "Nivel 2", "Nivel 3", "Genérico"])
        ws.append(["Servicios", None, "", None])
        wb.save(tmp_path / "PIVOT_MAESTRO.xlsx")

        from utils.config import Config
        from core.context import PipelineContext
        from etapas.etapa1 import AuditorTaxonomia

        config = Config(env="testing", PIVOT_DIR=tmp_path)
        stage = AuditorTaxonomia()
        stage._context = PipelineContext(config=config)
        stage._logger = logger_mock

        valores = stage._cargar_valores_pivot()
        assert len(valores["Nivel 2"]) == 0
        assert len(valores["Generico"]) == 0


# ---------------------------------------------------------------------------
# Tests: _procesar_campos
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestProcesarCampos:

    def _valores_pivot_minimos(self):
        from utils import normalizar_texto
        return {
            "Nivel 1": {normalizar_texto("Servicios"), normalizar_texto("Obras")},
            "Nivel 2": {normalizar_texto("Tecnología"), normalizar_texto("Construcción")},
            "Nivel 3": {normalizar_texto("Consultoría"), normalizar_texto("Edificación")},
            "Generico": {normalizar_texto("Software"), normalizar_texto("Civil")},
        }

    def test_valor_nuevo_supera_umbral_detectado(self, tmp_path, logger_mock):
        stage = _stage_con_config(tmp_path, logger_mock)
        stage.umbral_alerta = 2
        stage.detectar_similares = False

        df = pd.DataFrame({
            "Nivel 1": ["Valor Nuevo Inexistente"] * 3,
            "Nivel 2": ["Tecnología"] * 3,
            "Nivel 3": ["Consultoría"] * 3,
            "Genérico": ["Software"] * 3,
        })

        hallazgos = stage._procesar_campos(df, self._valores_pivot_minimos())

        campos_nuevos = [h["campo"] for h in hallazgos["nuevos"]]
        assert "Nivel 1" in campos_nuevos

    def test_valor_conocido_no_aparece_en_hallazgos(self, tmp_path, logger_mock):
        stage = _stage_con_config(tmp_path, logger_mock)
        stage.umbral_alerta = 1
        stage.detectar_similares = False

        df = pd.DataFrame({
            "Nivel 1": ["Servicios"] * 3,
            "Nivel 2": ["Tecnología"] * 3,
            "Nivel 3": ["Consultoría"] * 3,
            "Genérico": ["Software"] * 3,
        })

        hallazgos = stage._procesar_campos(df, self._valores_pivot_minimos())

        assert hallazgos["nuevos"] == []
        assert hallazgos["similares"] == []

    def test_valor_raro_por_debajo_umbral_ignorado(self, tmp_path, logger_mock):
        stage = _stage_con_config(tmp_path, logger_mock)
        stage.umbral_alerta = 5
        stage.detectar_similares = False

        df = pd.DataFrame({
            "Nivel 1": ["ValorNuevoRaro"],  # Solo 1 ocurrencia
            "Nivel 2": ["Tecnología"],
            "Nivel 3": ["Consultoría"],
            "Genérico": ["Software"],
        })

        hallazgos = stage._procesar_campos(df, self._valores_pivot_minimos())
        assert hallazgos["nuevos"] == []

    def test_columna_faltante_omitida_sin_error(self, tmp_path, logger_mock):
        stage = _stage_con_config(tmp_path, logger_mock)
        stage.umbral_alerta = 1
        stage.detectar_similares = False

        # Solo 'Nivel 1', sin las demás
        df = pd.DataFrame({"Nivel 1": ["ValorNuevo"] * 3})
        valores_pivot = {
            "Nivel 1": set(), "Nivel 2": set(), "Nivel 3": set(), "Generico": set(),
        }

        hallazgos = stage._procesar_campos(df, valores_pivot)

        assert isinstance(hallazgos, dict)
        assert "nuevos" in hallazgos and "similares" in hallazgos
        # Only Nivel 1 should have findings (others were skipped)
        campos_nuevos = {h["campo"] for h in hallazgos["nuevos"]}
        assert campos_nuevos <= {"Nivel 1"}

    def test_detectar_similares_no_crash(self, tmp_path, logger_mock):
        """Con detectar_similares=True, el procesamiento no falla."""
        from utils import normalizar_texto
        stage = _stage_con_config(tmp_path, logger_mock)
        stage.umbral_alerta = 1
        stage.detectar_similares = True
        stage.umbral_similitud = 0.7

        df = pd.DataFrame({
            "Nivel 1": ["Srvicios"] * 3,  # Typo
            "Nivel 2": ["Tecnología"] * 3,
            "Nivel 3": ["Consultoría"] * 3,
            "Genérico": ["Software"] * 3,
        })
        valores_pivot = {
            "Nivel 1": {"SERVICIOS"},
            "Nivel 2": {normalizar_texto("Tecnología")},
            "Nivel 3": {normalizar_texto("Consultoría")},
            "Generico": {normalizar_texto("Software")},
        }

        hallazgos = stage._procesar_campos(df, valores_pivot)
        assert "similares" in hallazgos

    def test_stats_actualizados_correctamente(self, tmp_path, logger_mock):
        stage = _stage_con_config(tmp_path, logger_mock)
        stage.umbral_alerta = 1
        stage.detectar_similares = False

        df = pd.DataFrame({
            "Nivel 1": ["NuevoA"] * 3,
            "Nivel 2": ["NuevoB"] * 3,
            "Nivel 3": ["Consultoría"] * 3,
            "Genérico": ["Software"] * 3,
        })
        from utils import normalizar_texto
        valores_pivot = {
            "Nivel 1": set(),
            "Nivel 2": set(),
            "Nivel 3": {normalizar_texto("Consultoría")},
            "Generico": {normalizar_texto("Software")},
        }

        stage._procesar_campos(df, valores_pivot)
        assert stage.stats["campos"] == 4  # All 4 fields processed
        assert stage.stats["nuevos"] == 2  # NuevoA and NuevoB


# ---------------------------------------------------------------------------
# Tests: _generar_reporte
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestGenerarReporte:

    def test_crea_archivo_excel(self, tmp_path, logger_mock):
        stage = _stage_con_config(tmp_path, logger_mock)
        hallazgos = {
            "nuevos": [
                {"campo": "Nivel 1", "valor": "X", "normalizado": "X",
                 "ocurrencias": 4, "porcentaje": 8.0}
            ],
            "similares": [],
        }

        ruta = stage._generar_reporte(hallazgos)

        assert ruta.exists()
        assert ruta.suffix == ".xlsx"

    def test_sheet_valores_nuevos_presente(self, tmp_path, logger_mock):
        stage = _stage_con_config(tmp_path, logger_mock)
        hallazgos = {
            "nuevos": [
                {"campo": "Nivel 1", "valor": "X", "normalizado": "X",
                 "ocurrencias": 4, "porcentaje": 8.0}
            ],
            "similares": [],
        }

        ruta = stage._generar_reporte(hallazgos)
        wb = load_workbook(ruta)
        assert "Valores Nuevos" in wb.sheetnames

    def test_sheet_valores_similares_presente(self, tmp_path, logger_mock):
        stage = _stage_con_config(tmp_path, logger_mock)
        hallazgos = {
            "nuevos": [],
            "similares": [
                {
                    "campo": "Nivel 1",
                    "valor_nuevo": "Servicvo",
                    "similares": [("SERVICIO", 0.85)],
                    "ocurrencias": 4,
                }
            ],
        }

        ruta = stage._generar_reporte(hallazgos)
        wb = load_workbook(ruta)
        assert "Valores Similares" in wb.sheetnames

    def test_nuevos_ordenados_descendente_por_ocurrencias(self, tmp_path, logger_mock):
        stage = _stage_con_config(tmp_path, logger_mock)
        hallazgos = {
            "nuevos": [
                {"campo": "Nivel 1", "valor": "Poco", "normalizado": "POCO",
                 "ocurrencias": 2, "porcentaje": 4.0},
                {"campo": "Nivel 1", "valor": "Muchos", "normalizado": "MUCHOS",
                 "ocurrencias": 10, "porcentaje": 20.0},
            ],
            "similares": [],
        }

        ruta = stage._generar_reporte(hallazgos)
        df = pd.read_excel(ruta, sheet_name="Valores Nuevos", engine="openpyxl")

        # Descending by ocurrencias within same campo
        assert df.iloc[0]["valor"] == "Muchos"
        assert df.iloc[1]["valor"] == "Poco"


# ---------------------------------------------------------------------------
# Tests: _cargar_parametros_pivot
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestCargarParametrosPivot:

    def test_carga_todos_los_parametros(self, etapa1):
        with patch("utils.config.Config.cargar_desde_pivot",
                   return_value={"Umbral de Alerta": "7", "Detectar Similares": "true",
                                 "Umbral Similitud": "0.92"}):
            etapa1._cargar_parametros_pivot()

        assert etapa1.umbral_alerta == 7
        assert etapa1.detectar_similares is True
        assert etapa1.umbral_similitud == pytest.approx(0.92)

    def test_detectar_similares_false(self, etapa1):
        with patch("utils.config.Config.cargar_desde_pivot",
                   return_value={"Detectar Similares": "false"}):
            etapa1._cargar_parametros_pivot()

        assert etapa1.detectar_similares is False

    def test_fallo_no_propaga_excepcion(self, etapa1):
        with patch("utils.config.Config.cargar_desde_pivot",
                   side_effect=Exception("PIVOT no disponible")):
            etapa1._cargar_parametros_pivot()  # No debe lanzar

        # Defaults conservados
        assert etapa1.umbral_alerta == 3

    def test_parametros_parciales_no_sobre_escriben_otros(self, etapa1):
        etapa1.umbral_alerta = 3  # Reset to default
        with patch("utils.config.Config.cargar_desde_pivot",
                   return_value={"Umbral de Alerta": "9"}):
            etapa1._cargar_parametros_pivot()

        assert etapa1.umbral_alerta == 9
        assert etapa1.detectar_similares is True  # No cambió


# ---------------------------------------------------------------------------
# Tests: run() — flujo completo
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestRun:

    def test_run_retorna_failure_si_pivot_no_existe(self, tmp_path, logger_mock):
        """Si PIVOT_MAESTRO no existe debe retornar StageResult(success=False)."""
        from utils.config import Config
        from core.context import PipelineContext
        from etapas.etapa1 import AuditorTaxonomia

        input_dir = tmp_path / "INPUT"
        input_dir.mkdir()
        licitaciones = input_dir / "Licitacion_Publicada.xlsx"
        _make_licitaciones_xlsx(licitaciones)

        # PIVOT_DIR no tiene PIVOT_MAESTRO.xlsx
        pivot_dir = tmp_path / "pivot_vacio"
        pivot_dir.mkdir()

        config = Config(
            env="testing",
            PIVOT_DIR=pivot_dir,
            INPUT_DIR=input_dir,
            HALLAZGOS_DIR=tmp_path / "hallazgos",
        )
        ctx = PipelineContext(config=config)
        stage = AuditorTaxonomia()

        result = stage.run(ctx)

        assert result.success is False
        assert result.error_message is not None

    def test_run_exitoso_sin_hallazgos(self, tmp_path, logger_mock):
        """Con todos los valores conocidos, run() tiene éxito y no genera archivo."""
        from utils.config import Config
        from core.context import PipelineContext
        from etapas.etapa1 import AuditorTaxonomia

        input_dir = tmp_path / "INPUT"
        input_dir.mkdir()
        hallazgos_dir = tmp_path / "hallazgos"
        hallazgos_dir.mkdir()
        licitaciones = input_dir / "Licitacion_Publicada.xlsx"
        _make_licitaciones_xlsx(licitaciones)  # Valores ya conocidos en el pivot
        _make_pivot_xlsx(tmp_path / "PIVOT_MAESTRO.xlsx")

        config = Config(
            env="testing",
            PIVOT_DIR=tmp_path,
            INPUT_DIR=input_dir,
            HALLAZGOS_DIR=hallazgos_dir,
            ETAPA1_UMBRAL_ALERTA=1,
        )
        ctx = PipelineContext(config=config)
        ctx.flags["allow_fallback"] = True
        stage = AuditorTaxonomia()

        with patch("utils.config.Config.cargar_desde_pivot", return_value={}):
            result = stage.run(ctx)

        assert result.success is True
        assert result.files_produced == []  # No hallazgos = no archivo

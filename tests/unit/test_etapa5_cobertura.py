"""
Tests adicionales para Etapa5 — cubre los huecos de cobertura restantes:
- _crear_reporte_inicial  → llama a guardar_formateado_reporte
- guardar_formateado_reporte (módulo compartido) → Excel con formato complejo
- _generar_analisis_cambios
- _guardar_sugerencias_pivot
- validate_inputs fallback
- run() success path
"""
import json
from unittest.mock import MagicMock, patch

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

# ---------------------------------------------------------------------------
# Fixtures locales compartidas
# ---------------------------------------------------------------------------

@pytest.fixture
def stage(config, logger_mock, tmp_path):
    """GeneradorReporteIncremental preparado con rutas temporales."""
    from core.context import PipelineContext
    from etapas.etapa5 import GeneradorReporteIncremental
    s = GeneradorReporteIncremental()
    ctx = PipelineContext(config=config)
    s._context = ctx
    s._logger = logger_mock
    s.dir_presentacion_incremental = tmp_path / "incremental"
    s.dir_presentacion_incremental.mkdir(parents=True)
    return s


@pytest.fixture
def df_minimo():
    return pd.DataFrame([
        {
            "Numero Adquisición": "2025-001",
            "Nombre": "Consultoría de TI",
            "Días para cierre": 10,
            "LINK": "https://www.mercadopublico.cl/xxx",
            "Región": "Metropolitana",
        }
    ])


@pytest.fixture
def df_con_todas_columnas():
    """DataFrame con todas las columnas formateadas por _guardar_formateado."""
    return pd.DataFrame([{
        "Numero Adquisición": "2025-TST",
        "Nombre": "N" * 60,
        "Descripción": "D" * 80,
        "Región": "Metropolitana",
        "Cliente (Organismo)": "Ministerio de Hacienda",
        "Fecha Publicación": "2026-01-01",
        "Hora Publicación": "08:00",
        "Fecha Inicio Preguntas": "2026-01-02",
        "Hora Inicio Preguntas": "08:00",
        "Fecha Cierre Preguntas": "2026-01-10",
        "Hora Cierre Preguntas": "17:00",
        "Fecha Apertura": "2026-01-15",
        "Hora Apertura": "09:00",
        "Fecha Cierre Licitación": "2026-01-20",
        "Hora Cierre Licitación": "17:00",
        "Fecha Adjudicación": "2026-02-01",
        "Hora Adjudicación": "12:00",
        "Monto Estimado (CLP)": 5000000,
        "Días para cierre": 8,
        "ONU": "Y",
        "Nivel 1": "Servicios de consultoría",
        "Nivel 2": "Tecnología de la información",
        "Nivel 3": "Desarrollo de software",
        "Genérico": "Consultoría TI",
        "Trazabilidad": "X",
        "LINK": "https://www.mercadopublico.cl/Procurement/Modules/RFB/DetailsAcquisition.aspx?idlicitacion=2025-TST",
    }])


@pytest.fixture
def reporte_anterior(tmp_path):
    ruta = tmp_path / "Reporte_Incremental_2025-01-01_00-00-00.xlsx"
    wb = Workbook()
    ws = wb.active  # type: ignore[union-attr]
    ws.title = "Vigentes"  # type: ignore[union-attr]
    ws.append(["Numero Adquisición", "Nombre", "Días para cierre"])  # type: ignore[union-attr]
    ws.append(["2025-001", "Existente", 20])  # type: ignore[union-attr]
    wb.save(ruta)
    return ruta


# ---------------------------------------------------------------------------
# _crear_reporte_inicial
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestCrearReporteInicial:

    def test_crea_excel_con_hoja_vigentes(self, stage, df_minimo):
        result = stage._crear_reporte_inicial(df_minimo)
        assert result['archivo_generado'].exists()
        assert result['tipo'] == 'inicial'
        wb = load_workbook(result['archivo_generado'])
        assert 'Vigentes' in wb.sheetnames

    def test_conteos_correctos(self, stage, df_minimo):
        result = stage._crear_reporte_inicial(df_minimo)
        assert result['licitaciones_nuevas'] == 1
        assert result['licitaciones_existentes'] == 0
        assert result['licitaciones_vencidas'] == 0

    def test_dataframe_vacio_no_crash(self, stage):
        result = stage._crear_reporte_inicial(pd.DataFrame())
        assert result['archivo_generado'].exists()
        assert result['licitaciones_nuevas'] == 0


# ---------------------------------------------------------------------------
# _guardar_formateado
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestGuardarFormateado:

    def test_genera_excel_con_encabezados(self, stage, df_con_todas_columnas, tmp_path):
        from utils.excel_formatter import guardar_formateado_reporte
        ruta = tmp_path / "salida.xlsx"
        guardar_formateado_reporte(df_con_todas_columnas, ruta, "Vigentes")
        assert ruta.exists()
        wb = load_workbook(ruta)
        ws = wb["Vigentes"]
        assert ws is not None
        assert ws.cell(1, 1).value == "Numero Adquisición"

    def test_congela_primera_fila(self, stage, df_minimo, tmp_path):
        from utils.excel_formatter import guardar_formateado_reporte
        ruta = tmp_path / "salida.xlsx"
        guardar_formateado_reporte(df_minimo, ruta, "Vigentes")
        wb = load_workbook(ruta)
        ws = wb["Vigentes"]
        assert ws is not None
        assert ws.freeze_panes == "A2"

    def test_link_convertido_en_hipervinculo(self, stage, df_con_todas_columnas, tmp_path):
        from utils.excel_formatter import guardar_formateado_reporte
        ruta = tmp_path / "salida.xlsx"
        guardar_formateado_reporte(df_con_todas_columnas, ruta, "Vigentes")
        wb = load_workbook(ruta)
        ws = wb["Vigentes"]
        assert ws is not None
        col_idx = list(df_con_todas_columnas.columns).index("LINK") + 1
        cell = ws.cell(2, col_idx)
        assert cell.hyperlink is not None or cell.value in ("Ver Licitación", df_con_todas_columnas["LINK"].iloc[0])

    def test_formato_condicional_dias_cierre(self, stage, df_con_todas_columnas, tmp_path):
        from utils.excel_formatter import guardar_formateado_reporte
        ruta = tmp_path / "salida.xlsx"
        guardar_formateado_reporte(df_con_todas_columnas, ruta, "Vigentes")
        wb = load_workbook(ruta)
        ws = wb["Vigentes"]
        assert ws is not None
        assert len(list(ws.conditional_formatting)) >= 1

    def test_altura_encabezado(self, stage, df_minimo, tmp_path):
        from utils.excel_formatter import guardar_formateado_reporte
        ruta = tmp_path / "salida.xlsx"
        guardar_formateado_reporte(df_minimo, ruta, "Vigentes")
        wb = load_workbook(ruta)
        ws = wb["Vigentes"]
        assert ws is not None
        assert ws.row_dimensions[1].height == 30

    def test_sin_columna_dias_cierre_no_crash(self, stage, tmp_path):
        from utils.excel_formatter import guardar_formateado_reporte
        df = pd.DataFrame([{"Nombre": "Consultoría", "Link": "http://x.cl"}])
        ruta = tmp_path / "salida.xlsx"
        guardar_formateado_reporte(df, ruta, "Vigentes")
        assert ruta.exists()

    def test_sin_columna_link_no_crash(self, stage, tmp_path):
        from utils.excel_formatter import guardar_formateado_reporte
        df = pd.DataFrame([{"Nombre": "Consultoría", "Días para cierre": 5}])
        ruta = tmp_path / "salida.xlsx"
        guardar_formateado_reporte(df, ruta, "Vigentes")
        assert ruta.exists()

    def test_alineacion_columnas_largas(self, stage, df_con_todas_columnas, tmp_path):
        from utils.excel_formatter import guardar_formateado_reporte
        ruta = tmp_path / "salida.xlsx"
        guardar_formateado_reporte(df_con_todas_columnas, ruta, "Vigentes")
        wb = load_workbook(ruta)
        ws = wb["Vigentes"]
        assert ws is not None
        col_idx = list(df_con_todas_columnas.columns).index("Nombre") + 1
        cell = ws.cell(2, col_idx)
        assert cell.alignment.wrap_text is True

    def test_formato_monto_numero(self, stage, df_con_todas_columnas, tmp_path):
        from utils.excel_formatter import guardar_formateado_reporte
        ruta = tmp_path / "salida.xlsx"
        guardar_formateado_reporte(df_con_todas_columnas, ruta, "Vigentes")
        wb = load_workbook(ruta)
        ws = wb["Vigentes"]
        assert ws is not None
        col_idx = list(df_con_todas_columnas.columns).index("Monto Estimado (CLP)") + 1
        cell = ws.cell(2, col_idx)
        assert '#' in cell.number_format


# ---------------------------------------------------------------------------
# _generar_analisis_cambios
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestGenerarAnalisisCambios:

    def test_sin_reporte_anterior(self, stage, df_minimo):
        mock_analisis_tax = {"categorias": []}
        mock_reporte = {"resumen": "ok"}
        with patch.object(stage.analizador, 'analizar_cambios_taxonomia', return_value=mock_analisis_tax), \
             patch.object(stage.analizador, 'generar_reporte_cambios', return_value=mock_reporte) as mock_gen:
            result = stage._generar_analisis_cambios(df_minimo, None)
        assert result == mock_reporte
        # generar_reporte_cambios debe recibir resultado sin reporte_anterior como arg
        mock_gen.assert_called_once()

    def test_con_reporte_anterior(self, stage, df_minimo, reporte_anterior):
        mock_analisis_tax = {"categorias": []}
        mock_analisis_rep = {"nuevas": [], "existentes": [], "vencidas": []}
        mock_reporte = {"resumen": "ok"}
        with patch.object(stage.analizador, 'analizar_cambios_taxonomia', return_value=mock_analisis_tax), \
             patch.object(stage.analizador, 'analizar_reporte_incremental', return_value=mock_analisis_rep), \
             patch.object(stage.analizador, 'generar_reporte_cambios', return_value=mock_reporte):
            result = stage._generar_analisis_cambios(df_minimo, reporte_anterior)
        assert result == mock_reporte

    def test_analisis_sin_reporte_genera_estadisticas_cero(self, stage, df_minimo):
        """Verifica que el fallback sin reporte anterior construye las estadísticas correctas."""
        captured = {}

        def mock_gen(analisis_tax, analisis_rep):
            captured['analisis_rep'] = analisis_rep
            return {}

        with patch.object(stage.analizador, 'analizar_cambios_taxonomia', return_value={}), \
             patch.object(stage.analizador, 'generar_reporte_cambios', side_effect=mock_gen):
            stage._generar_analisis_cambios(df_minimo, None)

        est = captured['analisis_rep']['estadisticas']
        assert est['existentes'] == 0
        assert est['vencidas'] == 0
        assert est['nuevas'] == len(df_minimo)


# ---------------------------------------------------------------------------
# _guardar_sugerencias_pivot
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestGuardarSugerenciasPivot:

    def test_crea_archivo_json(self, stage, config, tmp_path, monkeypatch):
        monkeypatch.setattr(config, 'LOG_DIR', tmp_path)
        analisis = {"sugerencias": ["agregar X"], "total": 1}
        stage._guardar_sugerencias_pivot(analisis)
        archivos = list(tmp_path.glob("sugerencias_pivot_etapa5_*.json"))
        assert len(archivos) == 1

    def test_json_legible_y_correcto(self, stage, config, tmp_path, monkeypatch):
        monkeypatch.setattr(config, 'LOG_DIR', tmp_path)
        analisis = {"clave": "valor", "numero": 42}
        stage._guardar_sugerencias_pivot(analisis)
        archivos = list(tmp_path.glob("sugerencias_pivot_etapa5_*.json"))
        assert len(archivos) == 1
        with archivos[0].open(encoding="utf-8") as f:
            data = json.load(f)
        assert data["clave"] == "valor"
        assert data["numero"] == 42


# ---------------------------------------------------------------------------
# validate_inputs — ramas de fallback
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestValidateInputsFallback:

    def test_sin_artefacto_sin_fallback_lanza_error(self, config):
        from core.context import PipelineContext
        from etapas.etapa5 import GeneradorReporteIncremental
        stage = GeneradorReporteIncremental()
        ctx = PipelineContext(config=config)
        with pytest.raises(ValueError, match="Falta artefacto"):
            stage.validate_inputs(ctx)

    def test_sin_artefacto_con_fallback_y_sin_archivos_lanza_error(self, config, tmp_path, monkeypatch):
        from core.context import PipelineContext
        from etapas.etapa5 import GeneradorReporteIncremental
        stage = GeneradorReporteIncremental()
        ctx = PipelineContext(config=config, flags={'allow_fallback': True})
        # Redirigir PRESENTACION_ORIGINAL_DIR a directorio vacío
        monkeypatch.setattr(config, 'PRESENTACION_ORIGINAL_DIR', tmp_path / "vacio")
        (tmp_path / "vacio").mkdir()
        with pytest.raises(FileNotFoundError):
            stage.validate_inputs(ctx)

    def test_archivo_entrada_no_existe_lanza_error(self, config, tmp_path):
        from core.context import PipelineContext
        from etapas.etapa5 import GeneradorReporteIncremental
        stage = GeneradorReporteIncremental()
        ruta_fantasma = tmp_path / "noexiste.xlsx"
        ctx = PipelineContext(config=config)
        ctx.add_artifact('etapa4_output', ruta_fantasma)
        with pytest.raises(FileNotFoundError):
            stage.validate_inputs(ctx)


# ---------------------------------------------------------------------------
# run() — camino de éxito
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestRun:

    def test_run_exito_con_artefacto(self, config, tmp_path, df_minimo, monkeypatch):
        from core.context import PipelineContext
        from etapas.etapa5 import GeneradorReporteIncremental
        stage = GeneradorReporteIncremental()
        stage._logger = MagicMock()
        stage._logger.finalize = MagicMock()
        stage._logger.section = MagicMock()
        stage._logger.info = MagicMock()
        stage._logger.error = MagicMock()

        incremental_dir = tmp_path / "incremental"
        incremental_dir.mkdir(parents=True)
        monkeypatch.setattr(config, 'PRESENTACION_INCREMENTAL_DIR', incremental_dir)

        entrada = tmp_path / "Reporte_Licitaciones_2026.xlsx"
        df_minimo.to_excel(entrada, index=False)

        ctx = PipelineContext(config=config)
        ctx.add_artifact('etapa4_output', entrada)

        mock_analisis = {"resumen": "ok", "sugerencias": []}
        with patch.object(stage.analizador, 'analizar_cambios_taxonomia', return_value={}), \
             patch.object(stage.analizador, 'generar_reporte_cambios', return_value=mock_analisis), \
             patch.object(stage, '_guardar_sugerencias_pivot'):
            result = stage.run(ctx)

        assert result.success is True
        assert result.stage_name == "incremental"

    def test_run_exito_sin_reporte_anterior_crea_inicial(self, config, tmp_path, df_minimo, monkeypatch):
        from core.context import PipelineContext
        from etapas.etapa5 import GeneradorReporteIncremental
        stage = GeneradorReporteIncremental()
        stage._logger = MagicMock()
        stage._logger.finalize = MagicMock()
        stage._logger.section = MagicMock()
        stage._logger.info = MagicMock()
        stage._logger.error = MagicMock()

        # Directorio incremental vacío → no habrá reporte anterior
        incremental_dir = tmp_path / "incremental_vacio"
        incremental_dir.mkdir(parents=True)
        monkeypatch.setattr(config, 'PRESENTACION_INCREMENTAL_DIR', incremental_dir)

        entrada = tmp_path / "Reporte_Licitaciones_2026.xlsx"
        df_minimo.to_excel(entrada, index=False)

        ctx = PipelineContext(config=config)
        ctx.add_artifact('etapa4_output', entrada)

        mock_analisis = {"resumen": "ok"}
        with patch.object(stage.analizador, 'analizar_cambios_taxonomia', return_value={}), \
             patch.object(stage.analizador, 'generar_reporte_cambios', return_value=mock_analisis), \
             patch.object(stage, '_guardar_sugerencias_pivot'):
            result = stage.run(ctx)

        assert result.success is True
        archivos = list(incremental_dir.glob("Reporte_Incremental_*.xlsx"))
        assert len(archivos) == 1

    def test_run_error_retorna_failure(self, config, tmp_path):
        from core.context import PipelineContext
        from etapas.etapa5 import GeneradorReporteIncremental
        stage = GeneradorReporteIncremental()
        stage._logger = MagicMock()
        stage._logger.finalize = MagicMock()
        ctx = PipelineContext(config=config)
        # Sin artefacto y sin fallback → validate_inputs lanzará ValueError
        result = stage.run(ctx)
        assert result.success is False
        assert result.error_message is not None

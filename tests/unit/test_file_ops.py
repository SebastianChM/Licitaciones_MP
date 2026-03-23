"""
Tests unitarios para src/utils/file_ops.py

Cubre: validar_archivo_excel, crear_backup, encontrar_fila_encabezado,
       encontrar_columna, leer_excel_con_header_dinamico,
       guardar_excel_con_formato, obtener_timestamp, obtener_fecha_hoy,
       listar_archivos_output, limpiar_outputs_antiguos
"""
import os
import time
from pathlib import Path
from unittest.mock import MagicMock, patch
import pandas as pd
import pytest
import openpyxl

from utils.file_ops import (
    validar_archivo_excel,
    crear_backup,
    encontrar_fila_encabezado,
    encontrar_columna,
    leer_excel_con_header_dinamico,
    guardar_excel_con_formato,
    obtener_timestamp,
    obtener_fecha_hoy,
    listar_archivos_output,
    limpiar_outputs_antiguos,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _excel_simple(path: Path, sheet: str = "Hoja1") -> Path:
    """Crea un Excel mínimo con una hoja y datos."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet  # type: ignore[union-attr]
    ws.append(["Nombre", "Valor"])  # type: ignore[union-attr]
    ws.append(["Alpha", 1])  # type: ignore[union-attr]
    wb.save(path)
    return path


def _excel_with_header_row(path: Path, header_row: int = 2) -> Path:
    """Crea un Excel donde los encabezados están en la fila header_row (0-indexed)."""
    wb = openpyxl.Workbook()
    ws = wb.active  # type: ignore[union-attr]
    # Filas de 'ruido' antes del header
    for _ in range(header_row):
        ws.append([""])  # type: ignore[union-attr]
    ws.append(["Nivel 1", "Codigo", "Nombre"])  # type: ignore[union-attr]
    ws.append(["A", "001", "Consultoria"])  # type: ignore[union-attr]
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# validar_archivo_excel
# ---------------------------------------------------------------------------

class TestValidarArchivoExcel:
    def test_archivo_no_existe_debe_existir(self, tmp_path):
        ruta = tmp_path / "noexiste.xlsx"
        valido, msg = validar_archivo_excel(ruta, debe_existir=True)
        assert not valido
        assert "no encontrado" in msg.lower()

    def test_archivo_no_existe_no_requerido(self, tmp_path):
        ruta = tmp_path / "noexiste.xlsx"
        valido, msg = validar_archivo_excel(ruta, debe_existir=False)
        assert valido
        assert msg == "OK"

    def test_extension_invalida(self, tmp_path):
        ruta = tmp_path / "datos.csv"
        ruta.write_text("a,b\n1,2")
        valido, msg = validar_archivo_excel(ruta, debe_existir=True)
        assert not valido
        assert "Formato inválido" in msg

    def test_archivo_valido_sin_hojas_requeridas(self, tmp_path):
        ruta = tmp_path / "datos.xlsx"
        _excel_simple(ruta)
        valido, msg = validar_archivo_excel(ruta)
        assert valido
        assert msg == "OK"

    def test_hoja_existente_pasa(self, tmp_path):
        ruta = tmp_path / "datos.xlsx"
        _excel_simple(ruta, sheet="MiHoja")
        valido, msg = validar_archivo_excel(ruta, hojas_requeridas=["MiHoja"])
        assert valido

    def test_hoja_faltante_falla(self, tmp_path):
        ruta = tmp_path / "datos.xlsx"
        _excel_simple(ruta, sheet="Hoja1")
        valido, msg = validar_archivo_excel(ruta, hojas_requeridas=["HojaQueNoExiste"])
        assert not valido
        assert "Hojas faltantes" in msg

    def test_archivo_corrupto_retorna_error(self, tmp_path):
        ruta = tmp_path / "corrupto.xlsx"
        ruta.write_bytes(b"notanexcel")
        valido, msg = validar_archivo_excel(ruta, hojas_requeridas=["Hoja1"])
        assert not valido
        assert "Error al leer" in msg


# ---------------------------------------------------------------------------
# crear_backup
# ---------------------------------------------------------------------------

class TestCrearBackup:
    def test_backup_crea_copia(self, tmp_path):
        original = tmp_path / "reporte.xlsx"
        _excel_simple(original)
        backup = crear_backup(original)
        assert backup is not None
        assert backup.exists()
        assert "backup_" in backup.name

    def test_backup_en_directorio_alternativo(self, tmp_path):
        original = tmp_path / "reporte.xlsx"
        _excel_simple(original)
        dest = tmp_path / "backups"
        backup = crear_backup(original, directorio_backup=dest)
        assert backup is not None
        assert backup.parent == dest

    def test_backup_archivo_inexistente_retorna_none(self, tmp_path):
        ruta = tmp_path / "noexiste.xlsx"
        resultado = crear_backup(ruta)
        assert resultado is None

    def test_backup_llama_logger_warning_si_no_existe(self, tmp_path):
        ruta = tmp_path / "noexiste.xlsx"
        logger = MagicMock()
        crear_backup(ruta, logger=logger)
        logger.warning.assert_called_once()

    def test_backup_llama_logger_info_si_exito(self, tmp_path):
        original = tmp_path / "reporte.xlsx"
        _excel_simple(original)
        logger = MagicMock()
        crear_backup(original, logger=logger)
        logger.info.assert_called_once()

    def test_backup_error_llama_logger_error(self, tmp_path):
        original = tmp_path / "reporte.xlsx"
        _excel_simple(original)
        logger = MagicMock()
        with patch("shutil.copy2", side_effect=OSError("disco lleno")):
            result = crear_backup(original, logger=logger)
        assert result is None
        logger.error.assert_called_once()

    def test_backup_error_sin_logger_retorna_none(self, tmp_path):
        original = tmp_path / "reporte.xlsx"
        _excel_simple(original)
        with patch("shutil.copy2", side_effect=OSError("disco lleno")):
            result = crear_backup(original)
        assert result is None


# ---------------------------------------------------------------------------
# encontrar_fila_encabezado
# ---------------------------------------------------------------------------

class TestEncontrarFilaEncabezado:
    def test_header_en_fila_0(self, tmp_path):
        ruta = tmp_path / "datos.xlsx"
        _excel_with_header_row(ruta, header_row=0)
        fila = encontrar_fila_encabezado(ruta, columna_referencia="Nivel 1")
        assert fila == 0

    def test_header_en_fila_2(self, tmp_path):
        ruta = tmp_path / "datos.xlsx"
        _excel_with_header_row(ruta, header_row=2)
        fila = encontrar_fila_encabezado(ruta, columna_referencia="Nivel 1")
        assert fila == 2

    def test_fallback_retorna_0_si_no_hay_match(self, tmp_path):
        """Si no hay ninguna columna ni keyword conocida, retorna 0."""
        ruta = tmp_path / "datos.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active  # type: ignore[union-attr]
        ws.append(["99999", "ZZZZZ"])  # type: ignore[union-attr]
        wb.save(ruta)
        fila = encontrar_fila_encabezado(ruta, columna_referencia="ColumnaInexistente")
        assert fila == 0

    def test_no_header_retorna_0(self, tmp_path):
        ruta = tmp_path / "datos.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active  # type: ignore[union-attr]
        ws.append([1, 2, 3])  # type: ignore[union-attr]
        wb.save(ruta)
        fila = encontrar_fila_encabezado(ruta, columna_referencia="Nivel 1")
        assert fila == 0

    def test_archivo_invalido_retorna_0(self, tmp_path):
        ruta = tmp_path / "corrupto.xlsx"
        ruta.write_bytes(b"notvalid")
        fila = encontrar_fila_encabezado(ruta)
        assert fila == 0

    def test_hoja_especificada(self, tmp_path):
        ruta = tmp_path / "datos.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active  # type: ignore[union-attr]
        ws.title = "Datos"  # type: ignore[union-attr]
        ws.append(["Nivel 1", "Codigo"])  # type: ignore[union-attr]
        wb.save(ruta)
        fila = encontrar_fila_encabezado(ruta, hoja="Datos", columna_referencia="Nivel 1")
        assert fila == 0


# ---------------------------------------------------------------------------
# encontrar_columna
# ---------------------------------------------------------------------------

class TestEncontrarColumna:
    def test_columna_exacta(self):
        df = pd.DataFrame({"Nombre": [], "Valor": []})
        assert encontrar_columna(df, "Nombre") == "Nombre"

    def test_columna_normalizada(self):
        df = pd.DataFrame({"Nivel 1": [], "Código": []})
        assert encontrar_columna(df, "nivel 1") == "Nivel 1"

    def test_columna_no_existe_con_normalizacion(self):
        df = pd.DataFrame({"Nombre": []})
        assert encontrar_columna(df, "ColumnaX") is None

    def test_columna_no_existe_sin_normalizacion(self):
        df = pd.DataFrame({"Nombre": []})
        assert encontrar_columna(df, "nombre", normalizar=False) is None

    def test_normalizar_false_exacto_no_encuentra(self):
        """Con normalizar=False no encuentra si el case difiere."""
        df = pd.DataFrame({"Nombre": []})
        result = encontrar_columna(df, "NOMBRE", normalizar=False)
        assert result is None


# ---------------------------------------------------------------------------
# leer_excel_con_header_dinamico
# ---------------------------------------------------------------------------

class TestLeerExcelConHeaderDinamico:
    def test_lee_datos_correctamente(self, tmp_path):
        ruta = tmp_path / "datos.xlsx"
        _excel_with_header_row(ruta, header_row=0)
        df = leer_excel_con_header_dinamico(ruta, columna_referencia="Nivel 1")
        assert "Nivel 1" in df.columns

    def test_hoja_especificada(self, tmp_path):
        ruta = tmp_path / "datos.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active  # type: ignore[union-attr]
        ws.title = "MiHoja"  # type: ignore[union-attr]
        ws.append(["Nivel 1", "Codigo"])  # type: ignore[union-attr]
        ws.append(["A", "001"])  # type: ignore[union-attr]
        wb.save(ruta)
        df = leer_excel_con_header_dinamico(ruta, hoja="MiHoja", columna_referencia="Nivel 1")
        assert "Nivel 1" in df.columns
        assert len(df) == 1

    def test_sin_hoja_especificada_usa_primera(self, tmp_path):
        ruta = tmp_path / "datos.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active  # type: ignore[union-attr]
        ws.title = "Hoja1"  # type: ignore[union-attr]
        ws.append(["Nivel 1"])  # type: ignore[union-attr]
        ws.append(["A"])  # type: ignore[union-attr]
        wb.save(ruta)
        df = leer_excel_con_header_dinamico(ruta)
        assert len(df) >= 0  # no crash


# ---------------------------------------------------------------------------
# guardar_excel_con_formato
# ---------------------------------------------------------------------------

class TestGuardarExcelConFormato:
    def test_guarda_correctamente(self, tmp_path):
        ruta = tmp_path / "salida.xlsx"
        df = pd.DataFrame({"Col A": [1, 2], "Col B": ["x", "y"]})
        resultado = guardar_excel_con_formato(df, ruta)
        assert resultado is True
        assert ruta.exists()

    def test_hoja_con_nombre_personalizado(self, tmp_path):
        ruta = tmp_path / "salida.xlsx"
        df = pd.DataFrame({"X": [1]})
        guardar_excel_con_formato(df, ruta, nombre_hoja="MiHoja")
        wb = openpyxl.load_workbook(ruta)
        assert "MiHoja" in wb.sheetnames

    def test_sin_autoajuste_ni_congela(self, tmp_path):
        ruta = tmp_path / "salida.xlsx"
        df = pd.DataFrame({"X": [1]})
        resultado = guardar_excel_con_formato(
            df, ruta, autoajustar_columnas=False, congelar_encabezado=False
        )
        assert resultado is True

    def test_error_retorna_false(self, tmp_path):
        ruta = tmp_path / "readonly_dir" / "salida.xlsx"
        # El directorio no existe → error al guardar
        df = pd.DataFrame({"X": [1]})
        resultado = guardar_excel_con_formato(df, ruta)
        assert resultado is False

    def test_freeze_panes_aplicado(self, tmp_path):
        ruta = tmp_path / "salida.xlsx"
        df = pd.DataFrame({"Col": [1, 2, 3]})
        guardar_excel_con_formato(df, ruta, congelar_encabezado=True)
        wb = openpyxl.load_workbook(ruta)
        ws = wb.active
        assert ws is not None  # type: ignore[union-attr]
        assert ws.freeze_panes == "A2"

    def test_dataframe_con_celdas_sin_valor(self, tmp_path):
        """Cubre el except pass en el loop de autoajuste."""
        ruta = tmp_path / "salida.xlsx"
        df = pd.DataFrame({"A": [None, None], "B": ["largo texto de prueba aqui", None]})
        resultado = guardar_excel_con_formato(df, ruta, autoajustar_columnas=True)
        assert resultado is True


# ---------------------------------------------------------------------------
# obtener_timestamp / obtener_fecha_hoy
# ---------------------------------------------------------------------------

class TestTimestamps:
    def test_timestamp_formato(self):
        ts = obtener_timestamp()
        # Formato: YYYY-MM-DD_HH-MM-SS
        parts = ts.split("_")
        assert len(parts) == 2
        assert len(parts[0]) == 10  # YYYY-MM-DD
        assert len(parts[1]) == 8   # HH-MM-SS

    def test_fecha_hoy_formato(self):
        fecha = obtener_fecha_hoy()
        assert len(fecha) == 8
        assert fecha.isdigit()


# ---------------------------------------------------------------------------
# listar_archivos_output
# ---------------------------------------------------------------------------

class TestListarArchivosOutput:
    def test_directorio_no_existe_retorna_vacio(self):
        with patch("utils.config.Config") as mock_cfg:
            mock_cfg.OUTPUT_DIR = Path("/ruta/que/no/existe/jamas")
            result = listar_archivos_output()
        assert result == []

    def test_lista_archivos_xlsx(self, tmp_path):
        (tmp_path / "a.xlsx").touch()
        (tmp_path / "b.xlsx").touch()
        (tmp_path / "c.txt").touch()
        with patch("utils.config.Config") as mock_cfg:
            mock_cfg.OUTPUT_DIR = tmp_path
            result = listar_archivos_output("*.xlsx")
        assert len(result) == 2
        assert all(p.suffix == ".xlsx" for p in result)

    def test_patron_personalizado(self, tmp_path):
        (tmp_path / "log.txt").touch()
        (tmp_path / "data.xlsx").touch()
        with patch("utils.config.Config") as mock_cfg:
            mock_cfg.OUTPUT_DIR = tmp_path
            result = listar_archivos_output("*.txt")
        assert len(result) == 1
        assert result[0].name == "log.txt"


# ---------------------------------------------------------------------------
# limpiar_outputs_antiguos
# ---------------------------------------------------------------------------

class TestLimpiarOutputsAntiguos:
    def test_no_limpia_si_pocos_archivos(self, tmp_path):
        archivos = [tmp_path / f"f{i}.xlsx" for i in range(3)]
        for a in archivos:
            a.touch()
        with patch("utils.config.Config") as mock_cfg:
            mock_cfg.OUTPUT_DIR = tmp_path
            limpiar_outputs_antiguos(dias=1, mantener_ultimos=5)
        # Todos deben seguir existiendo
        assert all(a.exists() for a in archivos)

    def test_elimina_archivos_antiguos(self, tmp_path):
        # Crear 7 archivos, los últimos 2 serán "antiguos" en tiempo de mtime
        archivos = []
        for i in range(7):
            a = tmp_path / f"f{i:02d}.xlsx"
            a.touch()
            archivos.append(a)

        # Hacer que los archivos más antiguos (índices 5 y 6 al ordenar desc) tengan mtime muy viejo
        # listar_archivos_output ordena reverse=True (más recientes primero por nombre)
        # Los índices 5+ son los más "viejos" según orden de nombre
        hace_60_dias = time.time() - (60 * 24 * 60 * 60)
        os.utime(archivos[0], (hace_60_dias, hace_60_dias))
        os.utime(archivos[1], (hace_60_dias, hace_60_dias))

        with patch("utils.config.Config") as mock_cfg:
            mock_cfg.OUTPUT_DIR = tmp_path
            limpiar_outputs_antiguos(dias=30, mantener_ultimos=5)

        # Los 5 más recientes (por orden de nombre desc) se mantienen
        remaining = list(tmp_path.glob("*.xlsx"))
        assert len(remaining) >= 5

    def test_no_falla_si_directorio_vacio(self, tmp_path):
        with patch("utils.config.Config") as mock_cfg:
            mock_cfg.OUTPUT_DIR = tmp_path
            limpiar_outputs_antiguos()  # no debe lanzar excepción

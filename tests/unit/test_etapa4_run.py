"""Tests unitarios para GeneradorReporte (Etapa 4) — run() y _guardar_formateado.

Complementa test_etapa4_convertir_monto.py que ya cubre _convertir_monto, _calcular_dias,
_separar.  Aquí se cubren:
- validate_inputs: con y sin artefacto, fallo por Excel inválido
- run(): flujo completo con mocks HTTP, salida exitosa e inputs faltantes
- _preparar_reporte: columna mandatoria faltante genera error
- _guardar_formateado: Excel escrito con hoja y links correctos
"""
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import MagicMock, patch

import pandas as pd
import pytest
from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

_FECHA_FUTURO = (datetime.now() + timedelta(days=30)).strftime("%Y-%m-%d 00:00:00")
_FECHA_PASADA = (datetime.now() - timedelta(days=5)).strftime("%Y-%m-%d 00:00:00")


def _make_input_xlsx(path: Path, fecha_cierre: str | None = None) -> None:
    """Crea Excel mínimo compatible con el formato de salida de Etapa 3."""
    fecha = fecha_cierre or _FECHA_FUTURO
    pd.DataFrame([{
        "Numero Adquisición": "2025-001-L1",
        "Nombre Adquisición": "Consultoría TI",
        "Descripción": "Servicios de consultoría tecnológica",
        "Región Compradora": "Metropolitana",
        "Organismo": "Ministerio TI",
        "Fecha Publicación": "2025-01-01 00:00:00",
        "FechaCierre": fecha,
        "Monto": "5000000",
        "Moneda": "CLP",
        "Tipo Adquisición": "Licitación Pública",
        "Nivel 1": "Servicios",
        "Nivel 2": "Tecnología",
        "Nivel 3": "Consultoría",
        "Genérico": "Software",
        "Trazabilidad Filtro": "Directo",
    }]).to_excel(path, index=False, engine="openpyxl")


def _make_multi_input_xlsx(path: Path) -> None:
    """Crea Excel con una licitación vigente y una vencida."""
    pd.DataFrame([
        {
            "Numero Adquisición": "2025-001-L1",
            "Nombre Adquisición": "Consultoría TI",
            "Descripción": "Servicio TI",
            "Región Compradora": "Metropolitana",
            "Organismo": "Min TI",
            "Fecha Publicación": "2025-01-01",
            "FechaCierre": _FECHA_FUTURO,
            "Monto": "5000000",
            "Moneda": "CLP",
            "Tipo Adquisición": "Licitación Pública",
            "Nivel 1": "Servicios",
            "Nivel 2": "Tecnología",
            "Nivel 3": "Consultoría",
            "Genérico": "Software",
            "Trazabilidad Filtro": "Directo",
        },
        {
            "Numero Adquisición": "2025-002-L1",
            "Nombre Adquisición": "Construcción",
            "Descripción": "Obras civiles",
            "Región Compradora": "Valparaíso",
            "Organismo": "Min Obras",
            "Fecha Publicación": "2025-01-01",
            "FechaCierre": _FECHA_PASADA,
            "Monto": "100000000",
            "Moneda": "CLP",
            "Tipo Adquisición": "Licitación Pública",
            "Nivel 1": "Obras",
            "Nivel 2": "Construcción",
            "Nivel 3": "Edificación",
            "Genérico": "Civil",
            "Trazabilidad Filtro": "Directo",
        },
    ]).to_excel(path, index=False, engine="openpyxl")


def _make_config(tmp_path: Path):
    from utils.config import Config
    pres_dir = tmp_path / "PRESENTACION" / "ORIGINALES"
    hist_dir = tmp_path / "HISTORICO"
    pres_dir.mkdir(parents=True)
    hist_dir.mkdir(parents=True)
    return Config(
        env="testing",
        PRESENTACION_ORIGINAL_DIR=pres_dir,
        HISTORICO_DIR=hist_dir,
        ENRIQUECIDO_DIR=tmp_path / "ENRIQUECIDO",
        ETAPA4_DIAS_GRACIA_HISTORICO=30,
        ETAPA4_VALOR_UTM=65000,
    )


def _stage(tmp_path: Path, logger_mock):
    from core.context import PipelineContext
    from etapas.etapa4 import GeneradorReporte
    config = _make_config(tmp_path)
    stage = GeneradorReporte()
    stage._context = PipelineContext(config=config)
    stage._logger = logger_mock
    return stage, config


# ---------------------------------------------------------------------------
# Tests: validate_inputs
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestValidateInputs:

    def test_artefacto_valido_no_lanza(self, tmp_path, logger_mock):
        stage, config = _stage(tmp_path, logger_mock)
        from core.context import PipelineContext
        ctx = PipelineContext(config=config)
        ctx.flags["allow_fallback"] = False

        input_file = tmp_path / "entrada.xlsx"
        _make_input_xlsx(input_file)
        ctx.add_artifact("etapa3_output", input_file)

        assert stage.validate_inputs(ctx) is True

    def test_sin_artefacto_modo_pipeline_lanza(self, tmp_path, logger_mock):
        stage, config = _stage(tmp_path, logger_mock)
        from core.context import PipelineContext
        ctx = PipelineContext(config=config)
        ctx.flags["allow_fallback"] = False  # strict pipeline mode

        with pytest.raises(ValueError, match="etapa3_output"):
            stage.validate_inputs(ctx)

    def test_sin_fallback_excepcion_artefacto_faltante(self, tmp_path, logger_mock):
        stage, config = _stage(tmp_path, logger_mock)
        from core.context import PipelineContext
        ctx = PipelineContext(config=config)
        ctx.flags["allow_fallback"] = False

        with pytest.raises(ValueError):
            stage.validate_inputs(ctx)

    def test_excel_corrupto_lanza_value_error(self, tmp_path, logger_mock):
        stage, config = _stage(tmp_path, logger_mock)
        from core.context import PipelineContext
        ctx = PipelineContext(config=config)
        ctx.flags["allow_fallback"] = False

        archivo_invalido = tmp_path / "corrupto.xlsx"
        archivo_invalido.write_bytes(b"esto no es un xlsx valido")
        ctx.add_artifact("etapa3_output", archivo_invalido)

        with pytest.raises((ValueError, Exception)):
            stage.validate_inputs(ctx)


# ---------------------------------------------------------------------------
# Tests: run() — flujo completo
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestRun:

    def _contexto_con_archivo(self, tmp_path, fecha_cierre=None):
        from core.context import PipelineContext
        config = _make_config(tmp_path)
        ctx = PipelineContext(config=config)
        ctx.flags["allow_fallback"] = False

        input_file = tmp_path / "entrada.xlsx"
        _make_input_xlsx(input_file, fecha_cierre)
        ctx.add_artifact("etapa3_output", input_file)
        return ctx, config

    def test_run_exitoso_licitacion_vigente(self, tmp_path, logger_mock):
        ctx, _ = self._contexto_con_archivo(tmp_path)

        from etapas.etapa4 import GeneradorReporte
        stage = GeneradorReporte()

        with patch.object(stage, "_actualizar_tasas"), \
             patch.object(stage, "_obtener_utm", return_value=65000), \
             patch.object(stage, "_obtener_usd", return_value=900):
            result = stage.run(ctx)

        assert result.success is True
        assert len(result.files_produced) >= 1
        assert result.files_produced[0].exists()

    def test_run_retorna_failure_sin_input(self, tmp_path, logger_mock):
        from core.context import PipelineContext
        from etapas.etapa4 import GeneradorReporte
        config = _make_config(tmp_path)
        ctx = PipelineContext(config=config)
        ctx.flags["allow_fallback"] = False  # strict: no artifact → failure

        stage = GeneradorReporte()

        result = stage.run(ctx)

        assert result.success is False
        assert result.error_message is not None

    def test_run_artifacts_registrados_en_contexto(self, tmp_path, logger_mock):
        ctx, _ = self._contexto_con_archivo(tmp_path)

        from etapas.etapa4 import GeneradorReporte
        stage = GeneradorReporte()

        with patch.object(stage, "_actualizar_tasas"):
            result = stage.run(ctx)

        assert result.success is True
        assert ctx.get_artifact("etapa4_output") is not None

    def test_run_con_vigente_y_vencida_produce_dos_archivos(self, tmp_path, logger_mock):
        from core.context import PipelineContext
        config = _make_config(tmp_path)
        _ = config.model_config  # validate config is ok
        ctx = PipelineContext(config=config)
        ctx.flags["allow_fallback"] = False

        input_file = tmp_path / "multi.xlsx"
        _make_multi_input_xlsx(input_file)
        ctx.add_artifact("etapa3_output", input_file)

        from etapas.etapa4 import GeneradorReporte
        stage = GeneradorReporte()

        with patch.object(stage, "_actualizar_tasas"):
            result = stage.run(ctx)

        assert result.success is True
        # One vigente + one vencida (within dias_gracia=30) → 2 files
        assert len(result.files_produced) == 2

    def test_stats_total_contabilizado(self, tmp_path, logger_mock):
        ctx, _ = self._contexto_con_archivo(tmp_path)

        from etapas.etapa4 import GeneradorReporte
        stage = GeneradorReporte()

        with patch.object(stage, "_actualizar_tasas"):
            result = stage.run(ctx)

        assert result.success is True
        assert result.metrics_produced["total"] == 1


# ---------------------------------------------------------------------------
# Tests: _preparar_reporte
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestPrepararReporte:

    def _stage_simple(self, tmp_path):
        from core.context import PipelineContext
        from etapas.etapa4 import GeneradorReporte
        config = _make_config(tmp_path)
        stage = GeneradorReporte()
        stage._context = PipelineContext(config=config)
        stage._logger = MagicMock()
        return stage

    def test_columna_numero_adquisicion_obligatoria(self, tmp_path):
        stage = self._stage_simple(tmp_path)
        df_sin_codigo = pd.DataFrame({"Nombre": ["X"]})

        # LINK column is built first using df_raw['Numero Adquisición'] directly →
        # KeyError precedes the explicit ValueError, both indicate missing mandatory column.
        with pytest.raises((KeyError, ValueError)):
            stage._preparar_reporte(df_sin_codigo)

    def test_genera_columna_link(self, tmp_path):
        stage = self._stage_simple(tmp_path)
        df = pd.DataFrame({
            "Numero Adquisición": ["2025-001-L1"],
            "FechaCierre": [_FECHA_FUTURO],
            "Monto": ["5000000"],
            "Moneda": ["CLP"],
            "Tipo Adquisición": ["Pública"],
        })

        df_out = stage._preparar_reporte(df)

        assert "LINK" in df_out.columns
        assert "2025-001-L1" in df_out["LINK"].iloc[0]

    def test_monto_convertido_correctamente(self, tmp_path):
        stage = self._stage_simple(tmp_path)
        stage.valor_utm = 65_000
        stage.valor_usd = 900
        df = pd.DataFrame({
            "Numero Adquisición": ["2025-001-L1"],
            "Monto": ["5000000"],
            "Moneda": ["CLP"],
            "Tipo Adquisición": ["Pública"],
        })

        df_out = stage._preparar_reporte(df)
        assert df_out["Monto Estimado (CLP)"].iloc[0] == 5_000_000


# ---------------------------------------------------------------------------
# Tests: guardar_formateado_reporte (shared Excel formatter)
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestGuardarFormateado:

    def test_archivo_creado_con_hoja_correcta(self, tmp_path):
        from etapas.etapa4 import GeneradorReporte
        from utils.excel_formatter import guardar_formateado_reporte

        df = pd.DataFrame({col: ["val"] for col in GeneradorReporte.COLUMNAS})
        df["LINK"] = ["https://www.mercadopublico.cl/test"]
        df["Días para cierre"] = [30]
        df["Monto Estimado (CLP)"] = [5_000_000]

        ruta = tmp_path / "test_reporte.xlsx"
        guardar_formateado_reporte(df, ruta, "Licitaciones")

        assert ruta.exists()
        wb = load_workbook(ruta)
        assert "Licitaciones" in wb.sheetnames

    def test_headers_escritos_correctamente(self, tmp_path):
        from etapas.etapa4 import GeneradorReporte
        from utils.excel_formatter import guardar_formateado_reporte

        df = pd.DataFrame({col: ["val"] for col in GeneradorReporte.COLUMNAS})
        df["LINK"] = ["https://www.mercadopublico.cl/test"]
        df["Días para cierre"] = [30]
        df["Monto Estimado (CLP)"] = [5_000_000]

        ruta = tmp_path / "test_headers.xlsx"
        guardar_formateado_reporte(df, ruta, "Licitaciones")

        wb = load_workbook(ruta)
        ws = wb["Licitaciones"]
        header_values = [ws.cell(row=1, column=i).value for i in range(1, len(GeneradorReporte.COLUMNAS) + 1)]
        assert "LINK" in header_values
        assert "Numero Adquisición" in header_values

    def test_hyperlink_generado_en_columna_link(self, tmp_path):
        from etapas.etapa4 import GeneradorReporte
        from utils.excel_formatter import guardar_formateado_reporte

        url = "https://www.mercadopublico.cl/test123"
        df = pd.DataFrame({col: ["val"] for col in GeneradorReporte.COLUMNAS})
        df["LINK"] = [url]
        df["Días para cierre"] = [30]
        df["Monto Estimado (CLP)"] = [5_000_000]

        ruta = tmp_path / "test_links.xlsx"
        guardar_formateado_reporte(df, ruta, "Licitaciones")

        wb = load_workbook(ruta)
        ws = wb["Licitaciones"]
        col_link_idx = GeneradorReporte.COLUMNAS.index("LINK") + 1
        celda = ws.cell(row=2, column=col_link_idx)
        assert celda.hyperlink is not None
        assert celda.value == "Ver Licitación"

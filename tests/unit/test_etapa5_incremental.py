"""Tests unitarios para GeneradorReporteIncremental (Etapa 5).

Cubre:
- _mover_licitaciones_vencidas: movimiento openpyxl entre hojas
- _actualizar_reporte_existente: copia base + actualizaciones
- _calcular_dias_cierre: lógica de días
- _generar_link_licitacion: generación de URL
"""
from unittest.mock import patch

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def etapa5(config, logger_mock):
    """GeneradorReporteIncremental con contexto y logger mockeados."""
    from core.context import PipelineContext
    from etapas.etapa5 import GeneradorReporteIncremental
    stage = GeneradorReporteIncremental()
    ctx = PipelineContext(config=config)
    stage._context = ctx
    stage._logger = logger_mock
    return stage


@pytest.fixture
def wb_vigentes():
    """Workbook openpyxl con hoja Vigentes y tres filas de datos."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Vigentes"
    ws.append(["Numero Adquisición", "Nombre", "Días para cierre"])
    ws.append(["2025-001", "Consultoría TI", 15])
    ws.append(["2025-002", "Maquinaria pesada", 5])
    ws.append(["2025-003", "Servicios digitales", 8])
    return wb


@pytest.fixture
def df_vencidas():
    """DataFrame con una licitación vencida (código 2025-001)."""
    return pd.DataFrame([{"Numero Adquisición": "2025-001", "Nombre": "Consultoría TI"}])


@pytest.fixture
def reporte_anterior_xlsx(tmp_path):
    """Excel temporal que simula el reporte anterior con hoja Vigentes."""
    ruta = tmp_path / "Reporte_Incremental_2025-01-01_00-00-00.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Vigentes"
    ws.append(["Numero Adquisición", "Nombre", "Días para cierre"])
    ws.append(["2025-001", "Existente", 20])
    wb.save(ruta)
    return ruta


@pytest.fixture
def df_datos_nuevos():
    return pd.DataFrame([
        {"Numero Adquisición": "2025-001", "Nombre": "Existente", "Días para cierre": 18},
        {"Numero Adquisición": "2025-NEW", "Nombre": "Nueva licitación", "Días para cierre": 30},
    ])


# ---------------------------------------------------------------------------
# _mover_licitaciones_vencidas
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestMoverLicitacionesVencidas:

    def test_crea_hoja_vencidas(self, etapa5, wb_vigentes, df_vencidas):
        ws = wb_vigentes["Vigentes"]
        etapa5._mover_licitaciones_vencidas(wb_vigentes, ws, df_vencidas, col_codigo_idx=1)
        assert "Vencidas" in wb_vigentes.sheetnames

    def test_copia_encabezado_a_vencidas(self, etapa5, wb_vigentes, df_vencidas):
        ws = wb_vigentes["Vigentes"]
        etapa5._mover_licitaciones_vencidas(wb_vigentes, ws, df_vencidas, col_codigo_idx=1)
        ws_v = wb_vigentes["Vencidas"]
        assert ws_v.cell(1, 1).value == "Numero Adquisición"

    def test_fila_vencida_aparece_en_hoja_vencidas(self, etapa5, wb_vigentes, df_vencidas):
        ws = wb_vigentes["Vigentes"]
        etapa5._mover_licitaciones_vencidas(wb_vigentes, ws, df_vencidas, col_codigo_idx=1)
        ws_v = wb_vigentes["Vencidas"]
        codigos = [ws_v.cell(r, 1).value for r in range(2, ws_v.max_row + 1)]
        assert "2025-001" in codigos

    def test_fila_vencida_eliminada_de_vigentes(self, etapa5, wb_vigentes, df_vencidas):
        ws = wb_vigentes["Vigentes"]
        etapa5._mover_licitaciones_vencidas(wb_vigentes, ws, df_vencidas, col_codigo_idx=1)
        codigos = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
        assert "2025-001" not in codigos

    def test_filas_no_vencidas_permanecen(self, etapa5, wb_vigentes, df_vencidas):
        ws = wb_vigentes["Vigentes"]
        etapa5._mover_licitaciones_vencidas(wb_vigentes, ws, df_vencidas, col_codigo_idx=1)
        codigos = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
        assert "2025-002" in codigos
        assert "2025-003" in codigos

    def test_df_vencidas_vacio_no_modifica_vigentes(self, etapa5, wb_vigentes):
        """Cuando vencidas está vacío, ninguna fila de Vigentes se elimina."""
        ws = wb_vigentes["Vigentes"]
        filas_antes = ws.max_row
        etapa5._mover_licitaciones_vencidas(wb_vigentes, ws, pd.DataFrame(), col_codigo_idx=1)
        # El número de filas en Vigentes no debe disminuir
        assert ws.max_row == filas_antes

    def test_columna_codigo_externo_reconocida(self, etapa5, wb_vigentes):
        ws = wb_vigentes["Vigentes"]
        vencidas = pd.DataFrame([{"Código Externo": "2025-002", "Nombre": "Maquinaria"}])
        etapa5._mover_licitaciones_vencidas(wb_vigentes, ws, vencidas, col_codigo_idx=1)
        ws_v = wb_vigentes["Vencidas"]
        codigos = [ws_v.cell(r, 1).value for r in range(2, ws_v.max_row + 1)]
        assert "2025-002" in codigos

    def test_sin_columna_codigo_loggea_warning(self, etapa5, wb_vigentes, logger_mock):
        ws = wb_vigentes["Vigentes"]
        vencidas = pd.DataFrame([{"OtraColumna": "valor"}])
        etapa5._mover_licitaciones_vencidas(wb_vigentes, ws, vencidas, col_codigo_idx=1)
        logger_mock.warning.assert_called_once()

    def test_sin_columna_codigo_no_modifica_vigentes(self, etapa5, wb_vigentes):
        ws = wb_vigentes["Vigentes"]
        filas_antes = ws.max_row
        vencidas = pd.DataFrame([{"OtraColumna": "valor"}])
        etapa5._mover_licitaciones_vencidas(wb_vigentes, ws, vencidas, col_codigo_idx=1)
        assert ws.max_row == filas_antes

    def test_hoja_vencidas_ya_existente_recibe_fila(self, etapa5, wb_vigentes, df_vencidas):
        """Si 'Vencidas' ya existe con encabezado, solo se agrega la fila."""
        wb = wb_vigentes
        ws_v_pre = wb.create_sheet("Vencidas")
        ws_v_pre.append(["Numero Adquisición", "Nombre", "Días para cierre"])
        ws = wb["Vigentes"]
        etapa5._mover_licitaciones_vencidas(wb, ws, df_vencidas, col_codigo_idx=1)
        ws_v = wb["Vencidas"]
        # Fila 1 = encabezado existente, fila 2 = dato copiado
        assert ws_v.cell(2, 1).value == "2025-001"



# ---------------------------------------------------------------------------
# _actualizar_reporte_existente
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestActualizarReporteExistente:

    def _analisis_mock(self, nuevas: pd.DataFrame, existentes: pd.DataFrame,
                       vencidas: pd.DataFrame) -> dict:
        return {
            'licitaciones_nuevas': nuevas,
            'licitaciones_existentes': existentes,
            'licitaciones_vencidas': vencidas,
        }

    def test_genera_archivo_xlsx(self, etapa5, tmp_path, reporte_anterior_xlsx, df_datos_nuevos):
        etapa5.dir_presentacion_incremental = tmp_path
        analisis = self._analisis_mock(
            df_datos_nuevos.iloc[[1]], df_datos_nuevos.iloc[[0]], pd.DataFrame()
        )
        with patch.object(etapa5.analizador, 'analizar_reporte_incremental', return_value=analisis):
            result = etapa5._actualizar_reporte_existente(df_datos_nuevos, reporte_anterior_xlsx)

        assert result['archivo_generado'].exists()
        assert result['archivo_generado'].suffix == ".xlsx"

    def test_conteos_en_resultado(self, etapa5, tmp_path, reporte_anterior_xlsx, df_datos_nuevos):
        etapa5.dir_presentacion_incremental = tmp_path
        analisis = self._analisis_mock(
            df_datos_nuevos.iloc[[1]], df_datos_nuevos.iloc[[0]], pd.DataFrame()
        )
        with patch.object(etapa5.analizador, 'analizar_reporte_incremental', return_value=analisis):
            result = etapa5._actualizar_reporte_existente(df_datos_nuevos, reporte_anterior_xlsx)

        assert result['licitaciones_nuevas'] == 1
        assert result['licitaciones_existentes'] == 1
        assert result['licitaciones_vencidas'] == 0
        assert result['tipo'] == 'incremental'

    def test_hoja_vigentes_preservada(self, etapa5, tmp_path, reporte_anterior_xlsx, df_datos_nuevos):
        etapa5.dir_presentacion_incremental = tmp_path
        analisis = self._analisis_mock(pd.DataFrame(), pd.DataFrame(), pd.DataFrame())
        with patch.object(etapa5.analizador, 'analizar_reporte_incremental', return_value=analisis):
            result = etapa5._actualizar_reporte_existente(df_datos_nuevos, reporte_anterior_xlsx)

        wb = load_workbook(result['archivo_generado'])
        assert 'Vigentes' in wb.sheetnames

    def test_fila_nueva_agregada_al_final(self, etapa5, tmp_path, reporte_anterior_xlsx, df_datos_nuevos):
        etapa5.dir_presentacion_incremental = tmp_path
        analisis = self._analisis_mock(
            df_datos_nuevos.iloc[[1]], pd.DataFrame(), pd.DataFrame()
        )
        with patch.object(etapa5.analizador, 'analizar_reporte_incremental', return_value=analisis):
            result = etapa5._actualizar_reporte_existente(df_datos_nuevos, reporte_anterior_xlsx)

        wb = load_workbook(result['archivo_generado'])
        ws = wb["Vigentes"]
        # Fila 1 = encabezado, fila 2 = dato base, fila 3 = nueva licitación
        codigos = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
        assert "2025-NEW" in codigos

    def test_vencidas_movidas_a_hoja_propia(self, etapa5, tmp_path, reporte_anterior_xlsx):
        etapa5.dir_presentacion_incremental = tmp_path
        # La licitación "2025-001" ya está en el reporte base y es marcada como vencida
        vencidas = pd.DataFrame([{"Numero Adquisición": "2025-001", "Nombre": "Existente",
                                   "Días para cierre": -5}])
        analisis = self._analisis_mock(pd.DataFrame(), pd.DataFrame(), vencidas)
        datos = pd.DataFrame([{"Numero Adquisición": "2025-001", "Días para cierre": -5}])
        with patch.object(etapa5.analizador, 'analizar_reporte_incremental', return_value=analisis):
            result = etapa5._actualizar_reporte_existente(datos, reporte_anterior_xlsx)

        wb = load_workbook(result['archivo_generado'])
        assert "Vencidas" in wb.sheetnames

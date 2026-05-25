"""Tests unitarios para core.profile_loader (ProfileRegistry, ProfileLoader, ProfileValidator)."""
import contextlib
from pathlib import Path

import openpyxl
import pytest

from core.filter_profile import EquipoInfo, FilterProfile
from core.profile_loader import (
    FILTROS_HEADER_ROW,
    HOJA_EQUIPOS,
    PivotConfigError,
    ProfileLoader,
    ProfileRegistry,
    ProfileValidator,
)

# ---------------------------------------------------------------------------
# Fixtures: construyen PIVOTs temporales con distintas estructuras
# ---------------------------------------------------------------------------

def _crear_pivot(
    path: Path,
    equipos: list[tuple[str, str, str, str, str]] | None = None,
    hojas_filtros: dict[str, list[list[str]]] | None = None,
    omitir_hoja_equipos: bool = False,
    columnas_equipos: list[str] | None = None,
) -> Path:
    """Crea un PIVOT_MAESTRO sintético.

    Args:
        equipos: filas para la hoja 00-Equipos (codigo, nombre, hoja, desc, activo).
        hojas_filtros: { nombre_hoja: [ [col0..col13] por fila de datos ] }
                       Se anteponen 4 filas de cabecera vacías + 1 header row.
        omitir_hoja_equipos: si True no crea la hoja 00-Equipos.
        columnas_equipos: nombres de columnas de 00-Equipos (default estándar).
    """
    wb = openpyxl.Workbook()
    default = wb.active
    assert default is not None

    if not omitir_hoja_equipos:
        ws_eq = wb.create_sheet(HOJA_EQUIPOS)
        cols = columnas_equipos or ["codigo", "nombre", "hoja_filtros", "descripcion", "activo"]
        ws_eq.append(cols)
        for fila in (equipos or []):
            ws_eq.append(list(fila))

    for hoja_nombre, filas_datos in (hojas_filtros or {}).items():
        ws = wb.create_sheet(hoja_nombre)
        # 4 filas de cabecera estilo PIVOT_MAESTRO
        for _ in range(FILTROS_HEADER_ROW):
            ws.append([""] * 14)
        # Header
        ws.append([
            "nombre", "nivel1", "nivel2", "nivel3",
            "nombre_exc", "nivel1_exc", "nivel2_exc", "nivel3_exc",
            "generico", "comp", "org", "valor",
            "bypass", "excl_dura",
        ])
        for fila in filas_datos:
            row = list(fila) + [""] * (14 - len(fila))
            ws.append(row)

    # Borrar la hoja default si nos sobra
    if default.title not in (wb.sheetnames[1:] if len(wb.sheetnames) > 1 else []):
        with contextlib.suppress(Exception):
            wb.remove(default)

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# ProfileRegistry
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestProfileRegistry:

    def test_lista_equipos_basicos(self, tmp_path):
        p = tmp_path / "PIVOT.xlsx"
        _crear_pivot(
            p,
            equipos=[
                ("TELECOM", "Telecom MP", "06-Telecom", "Telecom y redes", "TRUE"),
                ("ARQ",     "Arquitectura", "06-Arq",     "Edificios",       "TRUE"),
            ],
            hojas_filtros={},
        )
        reg = ProfileRegistry(p)
        equipos = reg.listar_equipos()
        assert len(equipos) == 2
        assert equipos[0].codigo == "TELECOM"
        assert equipos[1].nombre == "Arquitectura"

    def test_codigo_normalizado_a_mayuscula(self, tmp_path):
        p = tmp_path / "PIVOT.xlsx"
        _crear_pivot(p, equipos=[("telecom", "T", "06-T", "", "TRUE")], hojas_filtros={})
        reg = ProfileRegistry(p)
        assert reg.listar_equipos()[0].codigo == "TELECOM"

    def test_equipos_activos_filtra_inactivos(self, tmp_path):
        p = tmp_path / "PIVOT.xlsx"
        _crear_pivot(
            p,
            equipos=[
                ("A", "Equipo A", "06-A", "", "TRUE"),
                ("B", "Equipo B", "06-B", "", "FALSE"),
            ],
            hojas_filtros={},
        )
        reg = ProfileRegistry(p)
        activos = reg.equipos_activos()
        assert len(activos) == 1
        assert activos[0].codigo == "A"

    def test_buscar_codigo_existente(self, tmp_path):
        p = tmp_path / "PIVOT.xlsx"
        _crear_pivot(p, equipos=[("TELECOM", "T", "06-T", "", "TRUE")], hojas_filtros={})
        reg = ProfileRegistry(p)
        assert reg.buscar("TELECOM").codigo == "TELECOM"
        # Case insensitive
        assert reg.buscar("telecom").codigo == "TELECOM"

    def test_buscar_inexistente_lanza_error(self, tmp_path):
        p = tmp_path / "PIVOT.xlsx"
        _crear_pivot(p, equipos=[("A", "A", "06-A", "", "TRUE")], hojas_filtros={})
        reg = ProfileRegistry(p)
        with pytest.raises(PivotConfigError, match="no existe en el catálogo"):
            reg.buscar("FANTASMA")

    def test_pivot_inexistente_lanza_error(self, tmp_path):
        reg = ProfileRegistry(tmp_path / "no_existe.xlsx")
        with pytest.raises(PivotConfigError, match="no encontrado"):
            reg.listar_equipos()

    def test_hoja_equipos_faltante_lanza_error(self, tmp_path):
        p = tmp_path / "PIVOT.xlsx"
        _crear_pivot(p, omitir_hoja_equipos=True, hojas_filtros={"otra": []})
        reg = ProfileRegistry(p)
        with pytest.raises(PivotConfigError, match=HOJA_EQUIPOS):
            reg.listar_equipos()

    def test_columnas_requeridas_faltantes_lanza_error(self, tmp_path):
        p = tmp_path / "PIVOT.xlsx"
        _crear_pivot(
            p,
            equipos=[("X", "X", "06-X")],
            columnas_equipos=["codigo", "nombre", "otra"],  # falta hoja_filtros
            hojas_filtros={},
        )
        reg = ProfileRegistry(p)
        with pytest.raises(PivotConfigError, match="hoja_filtros"):
            reg.listar_equipos()


# ---------------------------------------------------------------------------
# ProfileLoader
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestProfileLoader:

    def test_carga_perfil_basico(self, tmp_path):
        p = tmp_path / "PIVOT.xlsx"
        _crear_pivot(
            p,
            equipos=[("TEL", "T", "06-T", "", "TRUE")],
            hojas_filtros={"06-T": [
                ["consultoria", "", "", "", "arriendo", "", "", "", "", "", "", "", "fibra optica", "salud"],
            ]},
        )
        equipo = ProfileRegistry(p).buscar("TEL")
        profile = ProfileLoader(p).cargar(equipo)

        assert profile.equipo.codigo == "TEL"
        # normalizar_texto uppercase + strip accents
        assert "CONSULTORIA" in profile.incluir["nombre"]
        assert "ARRIENDO" in profile.excluir["nombre"]
        assert "FIBRA OPTICA" in profile.bypass
        assert "SALUD" in profile.exclusion_dura

    def test_dedup_y_strip(self, tmp_path):
        p = tmp_path / "PIVOT.xlsx"
        _crear_pivot(
            p,
            equipos=[("X", "X", "06-X", "", "TRUE")],
            hojas_filtros={"06-X": [
                ["consultoria", "", "", "", "", "", "", "", "", "", "", "", "", ""],
                ["CONSULTORIA", "", "", "", "", "", "", "", "", "", "", "", "", ""],  # duplicado normalizado
                ["consultorIA", "", "", "", "", "", "", "", "", "", "", "", "", ""],  # otro duplicado
            ]},
        )
        equipo = ProfileRegistry(p).buscar("X")
        profile = ProfileLoader(p).cargar(equipo)
        assert list(profile.incluir["nombre"]).count("CONSULTORIA") == 1

    def test_hoja_inexistente_lanza_error(self, tmp_path):
        p = tmp_path / "PIVOT.xlsx"
        _crear_pivot(
            p,
            equipos=[("FANTASMA", "x", "06-NoExiste", "", "TRUE")],
            hojas_filtros={},
        )
        loader = ProfileLoader(p)
        equipo = ProfileRegistry(p).buscar("FANTASMA")
        with pytest.raises(PivotConfigError, match="06-NoExiste"):
            loader.cargar(equipo)

    def test_pivot_inexistente_lanza_error(self, tmp_path):
        loader = ProfileLoader(tmp_path / "no_existe.xlsx")
        equipo = EquipoInfo(codigo="X", nombre="x", hoja_filtros="06-X")
        with pytest.raises(PivotConfigError, match="no encontrado"):
            loader.cargar(equipo)


# ---------------------------------------------------------------------------
# ProfileValidator
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestProfileValidator:

    def _equipo(self) -> EquipoInfo:
        return EquipoInfo(codigo="T", nombre="T", hoja_filtros="06-T")

    def test_perfil_vacio_genera_advertencia(self):
        p = FilterProfile(equipo=self._equipo(), incluir={}, excluir={})
        problemas = ProfileValidator.validar(p)
        assert len(problemas) == 1
        assert "no tiene reglas de inclusión" in problemas[0]

    def test_perfil_completo_sin_problemas(self):
        p = FilterProfile(
            equipo=self._equipo(),
            incluir={"nombre": ["consultoria", "ingenieria"]},
            excluir={"nombre": ["arriendo"]},
            bypass=("fibra",),
        )
        assert ProfileValidator.validar(p) == []

    def test_keywords_de_un_caracter_detectadas(self):
        p = FilterProfile(
            equipo=self._equipo(),
            incluir={"nombre": ["a", "consultoria"]},
            excluir={},
        )
        problemas = ProfileValidator.validar(p)
        assert any("1 carácter" in pr for pr in problemas)

    def test_bypass_muy_corto_detectado(self):
        p = FilterProfile(
            equipo=self._equipo(),
            incluir={"nombre": ["consultoria"]},
            excluir={},
            bypass=("x",),
        )
        problemas = ProfileValidator.validar(p)
        assert any("Bypass" in pr for pr in problemas)


# ---------------------------------------------------------------------------
# ProfileCopier
# ---------------------------------------------------------------------------

from core.profile_loader import ProfileCopier


@pytest.mark.unit
class TestProfileCopier:
    """Onboarding: copiar reglas entre hojas-equipo."""

    def _pivot_dos_equipos(self, tmp_path: Path, destino_vacio: bool = True) -> Path:
        p = tmp_path / "PIVOT.xlsx"
        filas_dest = [] if destino_vacio else [["existente", "", "", "", "", "", "", "", "", "", "", "", "", ""]]
        _crear_pivot(
            p,
            equipos=[
                ("TELECOM", "Telecom",     "06-Tele", "", "TRUE"),
                ("ARQ",     "Arquitectura","06-Arq",  "", "TRUE"),
            ],
            hojas_filtros={
                "06-Tele": [
                    ["telecom", "redes", "fibra", "", "", "", "", "", "", "", "", "", "", ""],
                    ["antena", "", "", "", "", "", "", "", "", "", "", "", "", ""],
                ],
                "06-Arq": filas_dest,
            },
        )
        return p

    def test_copia_a_hoja_vacia(self, tmp_path):
        p = self._pivot_dos_equipos(tmp_path, destino_vacio=True)
        n = ProfileCopier(p).copiar("06-Tele", "06-Arq")
        assert n == 2
        # Verificar que las celdas se escribieron
        wb = openpyxl.load_workbook(p)
        ws = wb["06-Arq"]
        assert ws.cell(row=6, column=1).value == "telecom"
        assert ws.cell(row=7, column=1).value == "antena"
        wb.close()

    def test_falla_si_destino_tiene_datos_sin_sobrescribir(self, tmp_path):
        p = self._pivot_dos_equipos(tmp_path, destino_vacio=False)
        with pytest.raises(PivotConfigError, match="ya contiene reglas"):
            ProfileCopier(p).copiar("06-Tele", "06-Arq")

    def test_sobrescribir_explicito(self, tmp_path):
        p = self._pivot_dos_equipos(tmp_path, destino_vacio=False)
        n = ProfileCopier(p).copiar("06-Tele", "06-Arq", sobrescribir=True)
        assert n == 2
        wb = openpyxl.load_workbook(p)
        ws = wb["06-Arq"]
        assert ws.cell(row=6, column=1).value == "telecom"
        wb.close()

    def test_origen_igual_destino_falla(self, tmp_path):
        p = self._pivot_dos_equipos(tmp_path)
        with pytest.raises(PivotConfigError, match="no pueden ser la misma"):
            ProfileCopier(p).copiar("06-Tele", "06-Tele")

    def test_hoja_inexistente_falla(self, tmp_path):
        p = self._pivot_dos_equipos(tmp_path)
        with pytest.raises(PivotConfigError, match="no existe"):
            ProfileCopier(p).copiar("06-Tele", "06-Inexistente")

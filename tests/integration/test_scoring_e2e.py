"""Test de integración end-to-end: scoring en el flujo real del pipeline.

Verifica que cuando el pipeline ejecuta Etapa 4 sobre un archivo
producido por Etapa 2→3, el Excel resultante contiene la columna Score
con valores correctos, ordenados descendente y con formato condicional.

Ejecutar con:  pytest tests/integration/test_scoring_e2e.py -v -m integration
"""

from datetime import datetime, timedelta
from unittest.mock import patch

import pandas as pd
import pytest

pytestmark = pytest.mark.integration


def _crear_excel_etapa3(path, n=6):
    """Crea un Excel que simula la salida real de Etapa 3 (enriquecimiento).

    Incluye TODAS las columnas que el pipeline real produce:
    - Columnas base de Mercado Público (Etapa 0)
    - Columnas de filtrado (Etapa 2): Trazabilidad, Nivel Confianza,
      _n_matches_inclusion, _n_exclusiones_cercanas
    - Columnas de enriquecimiento (Etapa 3): API_*
    """
    ahora = datetime.now()
    filas = []
    perfiles = [
        # (nombre, descripcion, n_matches, n_exclusiones, nivel_confianza, monto, moneda, cierre_delta_dias)
        ("Consultoría TI infraestructura crítica", "Servicios profesionales de consultoría",
         4, 0, "ALTA", 500_000_000, "CLP", 10),
        ("Ingeniería telecomunicaciones red fibra", "Diseño y despliegue fibra óptica",
         5, 0, "ALTA", 1_500_000_000, "CLP", 30),
        ("Compra computadores escritorio", "Adquisición de equipos",
         1, 2, "BAJA", 3_000_000, "CLP", 1),
        ("Estudio geotécnico puente sector norte", "Estudio de suelos y geotecnia",
         3, 1, "ALTA", 100_000_000, "CLP", 5),
        ("Servicio aseo oficinas centrales", "Limpieza y mantención de oficinas",
         1, 4, "BAJA", 50_000_000, "CLP", -999),  # sin fecha (999 días)
        ("Diseño red eléctrica planta industrial", "Diseño de red eléctrica industrial",
         2, 3, "BAJA", 800_000_000, "CLP", -3),  # vencida
    ]

    for i, (nombre, desc, n_match, n_excl, confianza, monto, moneda, delta) in enumerate(perfiles):
        fecha_cierre = (ahora + timedelta(days=delta)).strftime("%Y-%m-%d %H:%M:%S") if delta != -999 else ""
        fecha_pub = (ahora - timedelta(days=5)).strftime("%Y-%m-%d %H:%M:%S")
        filas.append({
            "Numero Adquisición": f"2025-{100 + i}-L1",
            "Nombre Adquisición": nombre,
            "Descripción": desc,
            "Nivel 1": "Servicios",
            "Nivel 2": "Tecnología",
            "Nivel 3": "Consultoría",
            "Genérico": "Software",
            "Organismo": f"Organismo {i + 1}",
            "Tipo Adquisición": "Licitación Pública",
            "Descripción del producto/servicio": desc,
            "Monto": str(monto),
            "Moneda": moneda,
            "FechaPublicacion": fecha_pub,
            "FechaCierre": fecha_cierre,
            "CodigoProducto": f"ONU-{i}",
            # Columnas de Etapa 2 que sobreviven a Etapa 3
            "Trazabilidad Filtro": f"keyword{i} (Nombre Adquisición)" + (" | " + " | ".join(f"kw{k}" for k in range(1, n_match)) if n_match > 1 else ""),
            "Nivel Confianza": confianza,
            "_n_matches_inclusion": n_match,
            "_n_exclusiones_cercanas": n_excl,
            # Columnas API (Etapa 3)
            "API_Nombre": nombre,
            "API_Descripcion": desc,
            "API_Monto": str(monto),
            "API_Moneda": moneda,
            "API_FechaPublicacion": fecha_pub,
            "API_FechaCierre": fecha_cierre,
        })

    df = pd.DataFrame(filas)
    df.to_excel(path, index=False)
    return df


class TestScoringEndToEnd:
    """Verifica el scoring dentro del flujo real de Etapa 4."""

    def test_etapa4_produce_score_en_excel_real(self, context, tmp_path):
        """Etapa 4 ejecuta _execute completo y el Excel final tiene Score."""
        from etapas.etapa4 import GeneradorReporte

        # Crear input simulando salida real de Etapa 3
        archivo_e3 = tmp_path / "Licitaciones_Enriquecidas_test.xlsx"
        _crear_excel_etapa3(archivo_e3)

        # Configurar context como lo haría el pipeline real
        context.add_artifact("etapa3_output", archivo_e3)
        context.config.PRESENTACION_ORIGINAL_DIR = tmp_path / "presentacion"
        context.config.PRESENTACION_ORIGINAL_DIR.mkdir()
        context.config.HISTORICO_DIR = tmp_path / "historico"
        context.config.HISTORICO_DIR.mkdir()

        # Solo mockear las llamadas HTTP externas (UTM/USD)
        with patch.object(GeneradorReporte, '_actualizar_tasas'):
            etapa = GeneradorReporte()
            etapa.valor_utm = 65_000
            etapa.valor_usd = 950
            result = etapa.run(context)

        assert result.success is True, f"Etapa 4 falló: {result.error_message}"

        # Verificar que se generó el archivo
        archivo_salida = context.get_artifact("etapa4_output")
        assert archivo_salida is not None
        assert archivo_salida.exists()

        # Leer el Excel generado y verificar Score
        df_out = pd.read_excel(archivo_salida, engine="openpyxl")
        assert "Score" in df_out.columns, f"Columna Score ausente. Columnas: {list(df_out.columns)}"
        assert df_out.columns[0] == "Score", "Score debe ser la primera columna"

    def test_scores_ordenados_descendente(self, context, tmp_path):
        """Los scores están ordenados de mayor a menor en el Excel final."""
        from etapas.etapa4 import GeneradorReporte

        archivo_e3 = tmp_path / "Licitaciones_Enriquecidas_test.xlsx"
        _crear_excel_etapa3(archivo_e3)

        context.add_artifact("etapa3_output", archivo_e3)
        context.config.PRESENTACION_ORIGINAL_DIR = tmp_path / "presentacion"
        context.config.PRESENTACION_ORIGINAL_DIR.mkdir()
        context.config.HISTORICO_DIR = tmp_path / "historico"
        context.config.HISTORICO_DIR.mkdir()

        with patch.object(GeneradorReporte, '_actualizar_tasas'):
            etapa = GeneradorReporte()
            etapa.valor_utm = 65_000
            etapa.valor_usd = 950
            result = etapa.run(context)

        assert result.success is True
        df_out = pd.read_excel(context.get_artifact("etapa4_output"), engine="openpyxl")
        scores = df_out["Score"].tolist()
        assert scores == sorted(scores, reverse=True), f"Scores no están ordenados: {scores}"

    def test_scores_rango_valido(self, context, tmp_path):
        """Todos los scores están en el rango [1.0, 10.0]."""
        from etapas.etapa4 import GeneradorReporte

        archivo_e3 = tmp_path / "Licitaciones_Enriquecidas_test.xlsx"
        _crear_excel_etapa3(archivo_e3)

        context.add_artifact("etapa3_output", archivo_e3)
        context.config.PRESENTACION_ORIGINAL_DIR = tmp_path / "presentacion"
        context.config.PRESENTACION_ORIGINAL_DIR.mkdir()
        context.config.HISTORICO_DIR = tmp_path / "historico"
        context.config.HISTORICO_DIR.mkdir()

        with patch.object(GeneradorReporte, '_actualizar_tasas'):
            etapa = GeneradorReporte()
            etapa.valor_utm = 65_000
            etapa.valor_usd = 950
            result = etapa.run(context)

        assert result.success is True
        df_out = pd.read_excel(context.get_artifact("etapa4_output"), engine="openpyxl")
        for score in df_out["Score"]:
            assert 1.0 <= score <= 10.0, f"Score fuera de rango: {score}"

    def test_excel_tiene_formato_condicional_score(self, context, tmp_path):
        """El Excel final tiene reglas de formato condicional para Score (col A)."""
        from openpyxl import load_workbook

        from etapas.etapa4 import GeneradorReporte

        archivo_e3 = tmp_path / "Licitaciones_Enriquecidas_test.xlsx"
        _crear_excel_etapa3(archivo_e3)

        context.add_artifact("etapa3_output", archivo_e3)
        context.config.PRESENTACION_ORIGINAL_DIR = tmp_path / "presentacion"
        context.config.PRESENTACION_ORIGINAL_DIR.mkdir()
        context.config.HISTORICO_DIR = tmp_path / "historico"
        context.config.HISTORICO_DIR.mkdir()

        with patch.object(GeneradorReporte, '_actualizar_tasas'):
            etapa = GeneradorReporte()
            etapa.valor_utm = 65_000
            etapa.valor_usd = 950
            result = etapa.run(context)

        assert result.success is True
        wb = load_workbook(context.get_artifact("etapa4_output"))
        ws = wb.active

        # Buscar reglas de formato condicional aplicadas a columna A (Score)
        reglas_score = [
            rule
            for cf in ws.conditional_formatting._cf_rules
            if str(cf.sqref).startswith("A")
            for rule in cf.rules
        ]

        wb.close()
        assert len(reglas_score) == 4, f"Se esperaban 4 reglas para Score, hay {len(reglas_score)}"

    def test_score_con_un_decimal(self, context, tmp_path):
        """Cada score tiene exactamente un decimal."""
        from etapas.etapa4 import GeneradorReporte

        archivo_e3 = tmp_path / "Licitaciones_Enriquecidas_test.xlsx"
        _crear_excel_etapa3(archivo_e3)

        context.add_artifact("etapa3_output", archivo_e3)
        context.config.PRESENTACION_ORIGINAL_DIR = tmp_path / "presentacion"
        context.config.PRESENTACION_ORIGINAL_DIR.mkdir()
        context.config.HISTORICO_DIR = tmp_path / "historico"
        context.config.HISTORICO_DIR.mkdir()

        with patch.object(GeneradorReporte, '_actualizar_tasas'):
            etapa = GeneradorReporte()
            etapa.valor_utm = 65_000
            etapa.valor_usd = 950
            result = etapa.run(context)

        assert result.success is True
        df_out = pd.read_excel(context.get_artifact("etapa4_output"), engine="openpyxl")
        for score in df_out["Score"]:
            assert score == round(score, 1), f"Score sin 1 decimal: {score}"

    def test_score_formato_celda_excel(self, context, tmp_path):
        """Las celdas Score tienen formato numérico '0.0' y alineación center."""
        from openpyxl import load_workbook

        from etapas.etapa4 import GeneradorReporte

        archivo_e3 = tmp_path / "Licitaciones_Enriquecidas_test.xlsx"
        _crear_excel_etapa3(archivo_e3)

        context.add_artifact("etapa3_output", archivo_e3)
        context.config.PRESENTACION_ORIGINAL_DIR = tmp_path / "presentacion"
        context.config.PRESENTACION_ORIGINAL_DIR.mkdir()
        context.config.HISTORICO_DIR = tmp_path / "historico"
        context.config.HISTORICO_DIR.mkdir()

        with patch.object(GeneradorReporte, '_actualizar_tasas'):
            etapa = GeneradorReporte()
            etapa.valor_utm = 65_000
            etapa.valor_usd = 950
            result = etapa.run(context)

        assert result.success is True
        wb = load_workbook(context.get_artifact("etapa4_output"))
        ws = wb.active

        # Verificar header
        assert ws.cell(1, 1).value == "Score"

        # Verificar formato de celdas de datos
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row, 1)
            assert cell.number_format == "0.0", f"Fila {row}: formato={cell.number_format}"
            assert cell.alignment.horizontal == "center", f"Fila {row}: align={cell.alignment.horizontal}"

        wb.close()

    def test_vigentes_y_vencidas_tienen_score(self, context, tmp_path):
        """Tanto vigentes como vencidas tienen Score si DIAS_GRACIA permite histórico."""
        from etapas.etapa4 import GeneradorReporte

        archivo_e3 = tmp_path / "Licitaciones_Enriquecidas_test.xlsx"
        _crear_excel_etapa3(archivo_e3)

        context.add_artifact("etapa3_output", archivo_e3)
        context.config.PRESENTACION_ORIGINAL_DIR = tmp_path / "presentacion"
        context.config.PRESENTACION_ORIGINAL_DIR.mkdir()
        context.config.HISTORICO_DIR = tmp_path / "historico"
        context.config.HISTORICO_DIR.mkdir()

        with patch.object(GeneradorReporte, '_actualizar_tasas'):
            etapa = GeneradorReporte()
            etapa.valor_utm = 65_000
            etapa.valor_usd = 950
            etapa.dias_gracia = 30  # Incluir vencidas de hasta 30 días
            result = etapa.run(context)

        assert result.success is True
        # Verificar reporte principal (vigentes)
        df_vigentes = pd.read_excel(context.get_artifact("etapa4_output"), engine="openpyxl")
        assert "Score" in df_vigentes.columns
        assert all(1.0 <= s <= 10.0 for s in df_vigentes["Score"])

        # Verificar histórico (vencidas) si existe
        hist = context.get_artifact("etapa4_historico")
        if hist and hist.exists():
            df_vencidas = pd.read_excel(hist, engine="openpyxl")
            assert "Score" in df_vencidas.columns
            assert all(1.0 <= s <= 10.0 for s in df_vencidas["Score"])

    def test_scoring_sin_columnas_auxiliares_no_falla(self, context, tmp_path):
        """Si el input no tiene _n_matches/_n_exclusiones (legacy), scoring usa defaults."""
        from etapas.etapa4 import GeneradorReporte

        ahora = datetime.now()
        df_legacy = pd.DataFrame({
            "Numero Adquisición": ["LEG-001", "LEG-002"],
            "Nombre Adquisición": ["Consultoría TI", "Compra equipos"],
            "Monto": ["100000000", "5000000"],
            "Moneda": ["CLP", "CLP"],
            "FechaCierre": [
                (ahora + timedelta(days=10)).strftime("%Y-%m-%d"),
                (ahora + timedelta(days=3)).strftime("%Y-%m-%d"),
            ],
            "Nivel Confianza": ["ALTA", "BAJA"],
            "Trazabilidad Filtro": ["kw1 | kw2", "kw1"],
            # Sin _n_matches_inclusion ni _n_exclusiones_cercanas
        })
        archivo_legacy = tmp_path / "legacy.xlsx"
        df_legacy.to_excel(archivo_legacy, index=False)

        context.add_artifact("etapa3_output", archivo_legacy)
        context.config.PRESENTACION_ORIGINAL_DIR = tmp_path / "presentacion"
        context.config.PRESENTACION_ORIGINAL_DIR.mkdir()
        context.config.HISTORICO_DIR = tmp_path / "historico"
        context.config.HISTORICO_DIR.mkdir()

        with patch.object(GeneradorReporte, '_actualizar_tasas'):
            etapa = GeneradorReporte()
            etapa.valor_utm = 65_000
            etapa.valor_usd = 950
            result = etapa.run(context)

        assert result.success is True
        df_out = pd.read_excel(context.get_artifact("etapa4_output"), engine="openpyxl")
        assert "Score" in df_out.columns
        # Con 0 matches y 0 exclusiones (defaults), score es bajo pero válido
        for s in df_out["Score"]:
            assert 1.0 <= s <= 10.0

    def test_todas_vencidas_genera_excel_sin_crash(self, context, tmp_path):
        """Si todas las licitaciones están vencidas, el Excel se genera sin crash."""
        from etapas.etapa4 import GeneradorReporte

        ahora = datetime.now()
        df_vencidas = pd.DataFrame({
            "Numero Adquisición": ["VEN-001", "VEN-002"],
            "Nombre Adquisición": ["Licitación antigua 1", "Licitación antigua 2"],
            "Monto": ["50000000", "30000000"],
            "Moneda": ["CLP", "CLP"],
            "FechaCierre": [
                (ahora - timedelta(days=30)).strftime("%Y-%m-%d"),
                (ahora - timedelta(days=60)).strftime("%Y-%m-%d"),
            ],
            "Nivel Confianza": ["BAJA", "BAJA"],
            "_n_matches_inclusion": [1, 1],
            "_n_exclusiones_cercanas": [0, 0],
        })
        archivo = tmp_path / "todas_vencidas.xlsx"
        df_vencidas.to_excel(archivo, index=False)

        context.add_artifact("etapa3_output", archivo)
        context.config.PRESENTACION_ORIGINAL_DIR = tmp_path / "presentacion"
        context.config.PRESENTACION_ORIGINAL_DIR.mkdir()
        context.config.HISTORICO_DIR = tmp_path / "historico"
        context.config.HISTORICO_DIR.mkdir()

        with patch.object(GeneradorReporte, '_actualizar_tasas'):
            etapa = GeneradorReporte()
            etapa.valor_utm = 65_000
            etapa.valor_usd = 950
            etapa.dias_gracia = 0  # No incluir vencidas → 0 vigentes, 0 vencidas
            result = etapa.run(context)

        # Debe ser exitoso aunque ambos DataFrames estén vacíos
        assert result.success is True
        archivo_salida = context.get_artifact("etapa4_output")
        assert archivo_salida is not None
        assert archivo_salida.exists()
